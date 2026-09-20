import asyncio
import importlib.util
import json
import os
import sys
import tempfile
import threading
import time
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from jsonschema import Draft7Validator

ROOT = os.path.dirname(__file__)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from gateway.config import PlatformConfig
import plugins
import plugins.platforms

# Chạy test trực tiếp từ repo 2anh-zalo-bot nhưng dùng lõi Hermes đã cài trên máy.
# Ưu tiên hai plugin trong repo này, không vô tình test bản đã cài ở Hermes.
plugins.__path__ = [os.path.join(ROOT, "hermes-plugin"), *list(plugins.__path__)]
plugins.platforms.__path__ = [os.path.join(ROOT, "hermes-plugin"), *list(plugins.platforms.__path__)]
from plugins.platforms.zalo import adapter as zalo_adapter
from plugins.zalo_tools import tools as zalo_tools
from cron import jobs as real_cron_jobs

# Bốn thư viện sinh tệp cố tình KHÔNG có trong image hermes — xem
# docker/hermes-zalo.Dockerfile: "document capabilities were deferred". Đo trong
# container ngày 18/09/2026: thiếu cả bốn. Chúng được import muộn bên trong công
# cụ sinh tệp, nên chat vẫn chạy; nhưng bài kiểm thử nào dựng .docx thật sẽ đỏ ở
# đúng nơi duy nhất chạy được nó. Đỏ vì môi trường thiếu dép không phải đỏ vì mã
# sai, nên bỏ qua có nêu lý do — và khi nào cài dép, chúng tự chạy lại.
_FILE_MAKER_DEPS = ("docx", "fpdf", "openpyxl", "pptx")
_MISSING_FILE_MAKER_DEPS = [
    name for name in _FILE_MAKER_DEPS if importlib.util.find_spec(name) is None
]
requires_file_maker_deps = unittest.skipIf(
    bool(_MISSING_FILE_MAKER_DEPS),
    f"thiếu thư viện sinh tệp: {', '.join(_MISSING_FILE_MAKER_DEPS)}",
)


class DummyZaloTools:
    def set_turn_context(self, **kwargs):
        self.context = kwargs


class ZaloResolvedAllowlistTest(unittest.TestCase):
    def test_reads_guests_from_roster_without_returning_owners(self):
        owner_uid = "9000000000000000001"
        guest_uid = "9000000000000000002"
        with tempfile.TemporaryDirectory() as directory:
            roster_path = os.path.join(directory, "roster.json")
            with open(roster_path, "w", encoding="utf-8") as roster_file:
                json.dump({
                    "version": 1,
                    "owners": [owner_uid],
                    "guests": [guest_uid],
                }, roster_file)
            with patch.dict(os.environ, {"ZALO_ROSTER_FILE": roster_path}):
                adapter = zalo_adapter.ZaloAdapter(PlatformConfig(enabled=True, extra={}))
                self.assertEqual(adapter.resolved_allowlist_user_ids(), {guest_uid})

            os.unlink(roster_path)
            with patch.dict(os.environ, {"ZALO_ROSTER_FILE": roster_path}):
                self.assertEqual(adapter.resolved_allowlist_user_ids(), set())


class ZaloGuestTierSourcesTest(unittest.TestCase):
    """_is_guest phải nhìn cả hai nguồn khách, không chỉ env.

    Lỗi đo được trên VM 18/09/2026: một khách cấp bằng lệnh chat qua được cửa 1
    và qua được admission của gateway, rồi bị toolsets_for_source xếp là "không
    phải chủ nhân cũng không phải khách" và nhận TOOLSET_DENIED — không còn công
    cụ nào. Nhìn từ phía người dùng là bot không làm được gì, không phải một câu
    từ chối.
    """

    def test_guest_only_in_roster_is_still_a_guest(self):
        owner_uid = "9000000000000000001"
        chat_granted_guest = "9000000000000000002"
        with tempfile.TemporaryDirectory() as directory:
            roster_path = os.path.join(directory, "roster.json")
            with open(roster_path, "w", encoding="utf-8") as roster_file:
                json.dump({
                    "version": 1,
                    "owners": [owner_uid],
                    "guests": [chat_granted_guest],
                }, roster_file)
            # GATEWAY_ALLOWED_USERS rỗng: khách này CHỈ có trong roster, đúng
            # trạng thái sau một lần cấp bằng chat.
            with patch.dict(os.environ, {
                "ZALO_ROSTER_FILE": roster_path, "GATEWAY_ALLOWED_USERS": "",
            }):
                adapter = zalo_adapter.ZaloAdapter(PlatformConfig(enabled=True, extra={}))
                self.assertTrue(adapter._is_guest(chat_granted_guest))
                # Chủ nhân không được đi qua đường khách, và người lạ vẫn là người lạ.
                self.assertFalse(adapter._is_guest(owner_uid))
                self.assertFalse(adapter._is_guest("9000000000000000009"))


class CapturingSocket:
    def __init__(self):
        self.frames = []

    async def send(self, raw):
        self.frames.append(json.loads(raw))


class FakeToolContext:
    def __init__(self):
        self.handlers = {}
        self.hooks = {}

    def register_tool(self, **kwargs):
        self.handlers[kwargs["name"]] = kwargs["handler"]

    def register_hook(self, hook_name, callback):
        self.hooks.setdefault(hook_name, []).append(callback)


class FakeCronJobs:
    """Thay module cron.jobs của Hermes trong test — giữ job trong bộ nhớ."""

    parse_schedule = staticmethod(real_cron_jobs.parse_schedule)
    is_terminal_job = staticmethod(real_cron_jobs.is_terminal_job)
    effective_job_state = staticmethod(real_cron_jobs.effective_job_state)
    _ensure_croniter = staticmethod(real_cron_jobs._ensure_croniter)

    def __init__(self, jobs=None):
        self.jobs = [dict(job) for job in (jobs or [])]
        self.created = []
        self.removed = []

    @property
    def croniter(self):
        return real_cron_jobs.croniter

    def get_job(self, job_id):
        return next((job for job in self.jobs if job["id"] == job_id), None)

    def list_jobs(self, include_disabled=False):
        return [job for job in self.jobs if include_disabled or job.get("enabled", True)]

    def create_job(self, **kwargs):
        self.created.append(kwargs)
        job = {
            "id": f"new-{len(self.created)}", "enabled": True, "name": kwargs.get("name"),
            "schedule_display": kwargs["schedule"], "next_run_at": None,
            "prompt": kwargs["prompt"], "deliver": kwargs["deliver"], "origin": kwargs["origin"],
        }
        self.jobs.append(job)
        return job

    def remove_job(self, job_id):
        self.removed.append(job_id)
        before = len(self.jobs)
        self.jobs = [job for job in self.jobs if job["id"] != job_id]
        return len(self.jobs) < before


class ZaloAdapterMediaContextTest(unittest.IsolatedAsyncioTestCase):
    def make_adapter(self):
        adapter = zalo_adapter.ZaloAdapter(
            PlatformConfig(
                enabled=True,
                extra={
                    "bridge_url": "ws://127.0.0.1:9",
                    "reply_only_tagged": True,
                    "ack_gestures": False,
                },
            )
        )
        adapter._self_profile = {"user_id": "bot-uid", "display_name": "Lăng Tiêu"}
        adapter._flood.check = lambda _uid: None
        return adapter

    def test_bridge_url_contains_shared_token_without_changing_existing_query(self):
        self.assertEqual(
            zalo_adapter._authenticated_bridge_url(
                "ws://127.0.0.1:3873/path?existing=1", "bridge secret"
            ),
            "ws://127.0.0.1:3873/path?existing=1&token=bridge+secret",
        )
        with self.assertRaisesRegex(ValueError, "ZALO_BRIDGE_TOKEN"):
            zalo_adapter._authenticated_bridge_url("ws://127.0.0.1:3873", "")

    async def test_unmentioned_group_image_is_saved_as_context_only(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        await adapter._on_message(
            {
                "type": "message",
                "id": "m1",
                "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1",
                "senderName": "Yến",
                "text": "",
                "msgType": "chat.photo",
                "mediaUrls": ["https://example.com/photo.jpg"],
            }
        )

        self.assertEqual(handled, [])
        self.assertEqual(
            list(adapter._recent_group_messages["g1"])[0]["media_urls"],
            ["https://example.com/photo.jpg"],
        )

    async def test_later_mentioned_context_question_attaches_recent_image(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_cache(url):
            return f"C:/cache/{url.rsplit('/', 1)[-1]}"

        with patch.object(zalo_adapter, "cache_image_from_url", side_effect=fake_cache), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message(
                {
                    "type": "message",
                    "id": "m1",
                    "threadId": "g1",
                    "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                    "senderUid": "u1",
                    "senderName": "Yến",
                    "text": "",
                    "msgType": "chat.photo",
                    "mediaUrls": ["https://example.com/photo.jpg"],
                }
            )
            await adapter._on_message(
                {
                    "type": "message",
                    "id": "m2",
                    "threadId": "g1",
                    "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                    "senderUid": "u1",
                    "senderName": "Yến",
                    "text": "@Lăng Tiêu đây là xe máy điện hay xe đạp điện?",
                    "mentions": [{"uid": "bot-uid"}],
                }
            )

        self.assertEqual(len(handled), 1)
        event = handled[0]
        self.assertIn("đây là xe máy điện", event.text)
        self.assertEqual(event.media_urls, ["C:/cache/photo.jpg"])
        self.assertEqual(event.media_types, ["image/jpeg"])
        self.assertIn("Ngữ cảnh gần nhất trong nhóm Zalo", event.channel_context)
        self.assertIn("đã gửi 1 ảnh", event.channel_context)

    async def test_bare_tag_after_a_sticker_looks_at_what_was_just_sent(self):
        # Trong nhóm đệ ruột lúc 01:11 ngày 13/9: gửi sticker rồi tag trơ
        # "@Lăng Tiêu", bot hỏi lại "thầy cần gì ạ?" dù sticker nằm ngay trên.
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_cache(url):
            return "C:/cache/sticker.png"

        with patch.object(zalo_adapter, "cache_image_from_url", side_effect=fake_cache), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({
                "type": "message", "id": "s1", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Hải Anh",
                "text": "[Nhãn dán]", "msgType": "chat.sticker",
                "mediaUrls": ["https://zalo-api.zadn.vn/api/emoticon/sticker/webpc?eid=27703&size=130"],
                "attachments": [{
                    "url": "https://zalo-api.zadn.vn/api/emoticon/sticker/webpc?eid=27703&size=130",
                    "name": "sticker-27703.png", "mime": "image/png", "kind": "image",
                }],
            })
            await adapter._on_message({
                "type": "message", "id": "s2", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Hải Anh",
                "text": "@Lăng Tiêu", "mentions": [{"uid": "bot-uid"}],
            })

        self.assertEqual(len(handled), 1)
        self.assertEqual(handled[0].media_urls, ["C:/cache/sticker.png"])
        self.assertIn("Ngữ cảnh gần nhất trong nhóm Zalo", handled[0].channel_context)

    def test_owner_can_call_by_name_without_at_but_others_must_tag(self):
        # Anh Hải Anh gọi "Nhi ơi" trong nhóm mà bot không đáp (13–14/9). Chỉ chủ nhân
        # được gọi tên không cần @; người khác vẫn phải tag, kể cả tag tên ngắn.
        adapter = self.make_adapter()
        adapter._self_profile["display_name"] = "Uyển Nhi"

        self.assertTrue(adapter._is_mentioned({}, "Nhi ơi em cười vào mặt tiểu mi đi", is_owner=True))
        self.assertTrue(adapter._is_mentioned({}, "chào Nhi", is_owner=True))
        self.assertTrue(adapter._is_mentioned({}, "nhi oi", is_owner=True))
        self.assertFalse(adapter._is_mentioned({}, "Nhi ơi em cười vào mặt tiểu mi đi", is_owner=False))
        # Tag gõ tay, viết thường, hoặc chỉ tên ngắn — ai tag cũng được nhận.
        self.assertTrue(adapter._is_mentioned({}, "@uyển nhi đọc profile anh Khương", is_owner=False))
        self.assertTrue(adapter._is_mentioned({}, "@nhi giúp em với", is_owner=False))
        # Không nhầm chữ "nhi" nằm trong từ khác.
        self.assertFalse(adapter._is_mentioned({}, "bệnh nhi khoa đông quá", is_owner=False))

    def test_mention_only_recognizes_a_bare_call(self):
        adapter = self.make_adapter()
        self.assertTrue(adapter._mention_only("@Lăng Tiêu"))
        self.assertTrue(adapter._mention_only("  @Lăng Tiêu  !"))
        self.assertTrue(adapter._mention_only("@bot"))
        # Gọi tên kèm tiếng gọi vẫn là gọi suông.
        self.assertTrue(adapter._mention_only("@Lăng Tiêu ơi"))
        self.assertTrue(adapter._mention_only("Lăng Tiêu ơi"))
        self.assertTrue(adapter._mention_only("@Lăng Tiêu đâu rồi"))
        self.assertFalse(adapter._mention_only("@Lăng Tiêu soạn giúp anh thông báo"))
        self.assertFalse(adapter._mention_only("@Lăng Tiêu tóm tắt hộ anh"))
        self.assertFalse(adapter._mention_only(""))
        self.assertFalse(adapter._mention_only("chào cả nhà"))

    async def test_bare_call_replays_the_last_five_messages_of_the_group(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        with patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            for i, (who, line) in enumerate([
                ("Yến", "mai họp giao ban lúc mấy giờ ạ"),
                ("Trang", "8h nhé, phòng hội đồng"),
                ("Yến", "em xin phép đến muộn 15 phút"),
                ("Trang", "ok em, nhớ mang danh sách chi đoàn"),
                ("Giang", "danh sách em gửi trong nhóm hôm qua rồi ạ"),
                ("Yến", "vâng em xem lại"),
            ]):
                await adapter._on_message({
                    "type": "message", "id": f"c{i}", "threadId": "g1",
                    "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                    "senderUid": "u1", "senderName": who, "text": line,
                })
            await adapter._on_message({
                "type": "message", "id": "call", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Hải Anh",
                "text": "@Lăng Tiêu ơi", "mentions": [{"uid": "bot-uid"}],
            })

        self.assertEqual(len(handled), 1)
        context = handled[0].channel_context
        self.assertIn("Ngữ cảnh gần nhất trong nhóm Zalo", context)
        # Đúng 5 tin gần nhất trước câu gọi, không lấy tin thứ sáu.
        self.assertEqual(context.count("\n- "), 5)
        self.assertIn("danh sách em gửi trong nhóm hôm qua", context)
        self.assertNotIn("mai họp giao ban lúc mấy giờ", context)

    async def test_quote_image_is_attached_and_reply_context_set(self):
        adapter = self.make_adapter()
        handled = []
        dummy_tools = DummyZaloTools()

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_cache(url):
            return f"C:/cache/{url.rsplit('/', 1)[-1]}"

        with patch.object(zalo_adapter, "cache_image_from_url", side_effect=fake_cache), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=dummy_tools):
            await adapter._on_message(
                {
                    "type": "message",
                    "id": "m3",
                    "threadId": "g1",
                    "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                    "senderUid": "u1",
                    "senderName": "Yến",
                    "text": "@Lăng Tiêu cái này là gì?",
                    "mentions": [{"uid": "bot-uid"}],
                    "quote": {
                        "id": "q1",
                        "authorId": "bot-uid",
                        "authorName": "Lăng Tiêu",
                        "text": "",
                        "cliMsgId": "qc1",
                        "mediaUrls": ["https://example.com/quoted.jpg"],
                    },
                }
            )

        self.assertEqual(len(handled), 1)
        event = handled[0]
        self.assertEqual(event.media_urls, ["C:/cache/quoted.jpg"])
        self.assertEqual(event.reply_to_message_id, "q1")
        self.assertEqual(event.reply_to_text, "[Tin được reply có ảnh]")
        self.assertEqual(event.reply_to_author_id, "bot-uid")
        self.assertEqual(event.reply_to_author_name, "Lăng Tiêu")
        self.assertTrue(event.reply_to_is_own_message)
        self.assertEqual(dummy_tools.context["reply_msg_id"], "q1")
        self.assertEqual(dummy_tools.context["reply_cli_msg_id"], "qc1")
        self.assertTrue(dummy_tools.context["reply_is_own"])

    async def test_unreadable_quote_image_names_the_reason_instead_of_claiming_it_was_attached(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_cache(url):
            raise ValueError("Refusing to cache non-image data as .jpg (starts with: '�\\n')")

        with patch.object(zalo_adapter, "cache_image_from_url", side_effect=fake_cache), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message(
                {
                    "type": "message",
                    "id": "m4",
                    "threadId": "g1",
                    "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                    "senderUid": "u1",
                    "senderName": "Liên",
                    "text": "@Lăng Tiêu nhận xét ảnh này",
                    "mentions": [{"uid": "bot-uid"}],
                    "quote": {
                        "id": "q2", "authorId": "u1", "authorName": "Liên", "text": "",
                        "mediaUrls": ["https://photo-stal-17.zdn.vn/gr/heic/88b9/2aOb"],
                    },
                }
            )

        self.assertEqual(len(handled), 1)
        event = handled[0]
        self.assertEqual(event.media_urls, [])
        self.assertNotIn("đã được đính kèm", event.channel_context)
        self.assertIn("Không đọc được 1 ảnh", event.channel_context)
        self.assertIn("HEIC", event.channel_context)
        self.assertIn("gửi lại", event.channel_context)

    async def test_jxl_only_image_is_converted_to_jpeg_instead_of_failing(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_download(url):
            self.assertIn("/jxl/", url)
            return b"\xff\x0a du lieu jxl"

        with patch.object(zalo_adapter.ZaloAdapter, "_download_attachment", staticmethod(fake_download)), \
                patch.object(zalo_adapter.ZaloAdapter, "_jxl_to_jpeg", staticmethod(lambda data: b"\xff\xd8\xff jpeg")), \
                patch.object(zalo_adapter, "cache_image_from_bytes", return_value="C:/cache/converted.jpg"), \
                patch.object(zalo_adapter, "cache_image_from_url",
                             side_effect=AssertionError("ảnh JXL không được đi đường tải thẳng")), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({
                "type": "message", "id": "m5", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Liên",
                "text": "@Lăng Tiêu xem ảnh này", "mentions": [{"uid": "bot-uid"}],
                "msgType": "chat.photo",
                "mediaUrls": ["https://photo-stal-17.zdn.vn/gr/jxl/88b9/2aOb"],
            })

        self.assertEqual(len(handled), 1)
        self.assertEqual(handled[0].media_urls, ["C:/cache/converted.jpg"])
        self.assertNotIn("Không đọc được", handled[0].channel_context or "")

    async def test_jxl_image_without_decoder_says_what_is_missing(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        def missing_decoder(_data):
            raise RuntimeError(zalo_adapter._JXL_DECODER_MISSING)

        async def fake_download(url):
            return b"\xff\x0a du lieu jxl"

        with patch.object(zalo_adapter.ZaloAdapter, "_download_attachment", staticmethod(fake_download)), \
                patch.object(zalo_adapter.ZaloAdapter, "_jxl_to_jpeg", staticmethod(missing_decoder)), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({
                "type": "message", "id": "m6", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Liên",
                "text": "@Lăng Tiêu xem ảnh này", "mentions": [{"uid": "bot-uid"}],
                "msgType": "chat.photo",
                "mediaUrls": ["https://photo-stal-17.zdn.vn/gr/jxl/88b9/2aOb"],
            })

        self.assertEqual(len(handled), 1)
        event = handled[0]
        self.assertEqual(event.media_urls, [])
        self.assertIn("JPEG XL", event.channel_context)
        self.assertIn("pillow-jxl-plugin", event.channel_context)

    def test_jxl_to_jpeg_really_decodes_when_the_plugin_is_installed(self):
        try:
            import io
            import pillow_jxl  # noqa: F401
            from PIL import Image
        except ImportError:
            self.skipTest("máy này chưa cài pillow-jxl-plugin")

        buf = io.BytesIO()
        Image.new("RGB", (24, 16), (200, 30, 30)).save(buf, format="JXL")
        self.assertTrue(buf.getvalue().startswith(b"\xff\x0a"))

        jpeg = zalo_adapter.ZaloAdapter._jxl_to_jpeg(buf.getvalue())
        self.assertTrue(jpeg.startswith(b"\xff\xd8\xff"))
        with Image.open(io.BytesIO(jpeg)) as img:
            self.assertEqual((img.format, img.size), ("JPEG", (24, 16)))

    async def test_sticker_image_from_the_bridge_is_attached(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_cache(url):
            return f"C:/cache/{url.rsplit('/', 1)[-1]}"

        with patch.object(zalo_adapter, "cache_image_from_url", side_effect=fake_cache), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({
                "type": "message", "id": "st1", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Trang",
                "text": "[Nhãn dán: cười lăn]", "mentions": [{"uid": "bot-uid"}],
                "msgType": "chat.sticker",
                "mediaUrls": ["https://zalo.vn/sticker/4001.png"],
                "attachments": [{
                    "url": "https://zalo.vn/sticker/4001.png",
                    "name": "sticker-4001.png", "mime": "image/png", "kind": "image",
                }],
            })

        self.assertEqual(len(handled), 1)
        self.assertEqual(handled[0].media_urls, ["C:/cache/4001.png"])
        self.assertIn("[Nhãn dán: cười lăn]", handled[0].text)

    def test_link_cards_stay_blocked_while_classified_attachments_pass(self):
        # Danh sách loại trừ vẫn chặn việc đoán URL từ thẻ chia sẻ link…
        self.assertFalse(zalo_adapter._frame_carries_media(
            {"msgType": "chat.recommended", "mediaUrls": ["https://vt.tiktok.com/abc"]}))
        self.assertFalse(zalo_adapter._frame_carries_media(
            {"msgType": "chat.sticker", "attachments": []}))
        # …nhưng cầu nối đã phân loại sẵn thì đó là khẳng định, không phải đoán.
        self.assertTrue(zalo_adapter._frame_carries_media(
            {"msgType": "chat.sticker",
             "attachments": [{"url": "https://zalo.vn/sticker/1.png", "mime": "image/png"}]}))
        self.assertTrue(zalo_adapter._frame_carries_media({"msgType": "chat.photo"}))

    def test_is_jxl_reads_both_the_path_and_the_mime(self):
        self.assertTrue(zalo_adapter._is_jxl("https://photo-stal-17.zdn.vn/gr/jxl/88b9/2aOb"))
        self.assertTrue(zalo_adapter._is_jxl("https://x/anh.JXL"))
        self.assertTrue(zalo_adapter._is_jxl("https://x/anh", "image/jxl"))
        self.assertFalse(zalo_adapter._is_jxl("https://photo-stal-15.zdn.vn/gr/jpg/864/2aO.jpg"))
        self.assertFalse(zalo_adapter._is_jxl("https://x/jxl-tin-tuc/anh.jpg"))

    async def test_owner_only_group_answers_only_the_owner_but_keeps_everyone_as_context(self):
        # Nhóm cộng đồng 994 người: bot vào để nghe và tổng hợp, chỉ chủ nhân gọi được.
        adapter = self.make_adapter()
        adapter._owner_only_groups = {"g-cong-dong"}
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        def member_msg(msg_id, thread, uid, name, text, tag=True):
            frame = {
                "type": "message", "id": msg_id, "threadId": thread,
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": uid, "senderName": name, "text": text,
            }
            if tag:
                frame["mentions"] = [{"uid": "bot-uid"}]
            return frame

        with patch.object(adapter, "_is_owner", side_effect=lambda uid: uid == "owner-uid"), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message(member_msg("m1", "g-cong-dong", "u-la", "Người lạ",
                                                 "@Lăng Tiêu tóm tắt giúp em với"))
            await adapter._on_message(member_msg("m2", "g-cong-dong", "u-la2", "Người lạ 2",
                                                 "mọi người thấy agent này thế nào", tag=False))
            self.assertEqual(handled, [], "người lạ tag trong nhóm chỉ-chủ-nhân không được đánh thức bot")

            await adapter._on_message(member_msg("m3", "g-cong-dong", "owner-uid", "Hải Anh", "@Lăng Tiêu"))
            self.assertEqual(len(handled), 1)
            # Tin của người lạ vẫn nằm trong ngữ cảnh để chủ nhân gọi suông là đọc được.
            self.assertIn("mọi người thấy agent này thế nào", handled[0].channel_context)

            # Nhóm khác không bị ảnh hưởng.
            await adapter._on_message(member_msg("m4", "g-khac", "u-la", "Người lạ", "@Lăng Tiêu chào em"))
            self.assertEqual(len(handled), 2)

    def test_owner_only_groups_is_read_from_config_list_or_env_string(self):
        from gateway.config import PlatformConfig
        with patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            from_list = zalo_adapter.ZaloAdapter(PlatformConfig(
                enabled=True, extra={"owner_only_groups": ["39633293852382968"]}))
        self.assertEqual(from_list._owner_only_groups, {"39633293852382968"})
        with patch.dict(os.environ, {"ZALO_OWNER_ONLY_GROUPS": "111111111111111111, 222222222222222222"}), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            from_env = zalo_adapter.ZaloAdapter(PlatformConfig(enabled=True, extra={}))
        self.assertEqual(from_env._owner_only_groups, {"111111111111111111", "222222222222222222"})

    async def test_file_pulled_from_group_context_keeps_its_name_and_type(self):
        # Nhóm y tế 16:26 ngày 13/9: anh Trung gửi PDF (chưa tag bot), sau đó có
        # người tag bot nhờ đọc. Tệp móc từ ngữ cảnh mất tên và loại, URL không
        # đuôi nên bị đẩy vào đường ảnh: "Refusing to cache non-image data".
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle
        url = "https://file-stal-19.dlfl.vn/gr/4944b1f4cb33166d4f22/2aOboQyKriSQ4hH3h5SSgtGR0be"
        cached = []

        async def fake_download(_url):
            return b"%PDF-1.7 ban ve"

        async def fake_extract(_path):
            return "Bản vẽ bố trí khoa khám bệnh"

        def fake_cache(data, *, filename="", mime_type="", default_kind=None):
            cached.append((filename, mime_type))
            return SimpleNamespace(path=f"C:/cache/documents/{filename}", media_type=mime_type,
                                   kind="document", display_name=filename)

        with patch.object(zalo_adapter.ZaloAdapter, "_download_attachment", staticmethod(fake_download)), \
                patch.object(zalo_adapter.ZaloAdapter, "_document_text", staticmethod(fake_extract)), \
                patch.object(zalo_adapter, "cache_media_bytes", fake_cache), \
                patch.object(zalo_adapter, "cache_image_from_url",
                             side_effect=AssertionError("tệp PDF không được đi đường ảnh")), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({
                "type": "message", "id": "f1", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u2", "senderName": "Luong Dinh Trung",
                "text": "090926.pdf", "msgType": "share.file",
                "mediaUrls": [url],
                "attachments": [{"url": url, "name": "090926.pdf", "mime": "application/pdf", "kind": "document"}],
            })
            await adapter._on_message({
                "type": "message", "id": "t1", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Hải Anh",
                "text": "@Lăng Tiêu xem lại bản này nhé", "mentions": [{"uid": "bot-uid"}],
            })

        self.assertEqual(len(handled), 1)
        self.assertEqual(cached, [("090926.pdf", "application/pdf")])
        self.assertIn("Bản vẽ bố trí khoa khám bệnh", handled[0].text)
        self.assertNotIn("Không đọc được", handled[0].channel_context or "")

    async def test_pdf_attachment_becomes_a_document_not_an_image(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_download(url):
            return b"%PDF-1.7 noi dung"

        async def fake_extract(_path):
            return "Số: 21/KH-ĐTN\nKẾ HOẠCH tổ chức cuộc thi Tiếng nói xanh mùa IV"

        def fake_cache(data, *, filename="", mime_type="", default_kind=None):
            return SimpleNamespace(
                path=f"C:/cache/documents/doc_{filename}", media_type=mime_type or "application/pdf",
                kind="document", display_name=filename,
            )

        with patch.object(zalo_adapter.ZaloAdapter, "_download_attachment", staticmethod(fake_download)), \
                patch.object(zalo_adapter.ZaloAdapter, "_document_text", staticmethod(fake_extract)), \
                patch.object(zalo_adapter, "cache_media_bytes", fake_cache), \
                patch.object(zalo_adapter, "cache_image_from_url", side_effect=AssertionError("tệp không được đi đường ảnh")), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({
                "type": "message", "id": "f1", "threadId": "g1",
                "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "u1", "senderName": "Liên",
                "text": "@Lăng Tiêu trong file này có link ko?",
                "mentions": [{"uid": "bot-uid"}],
                "msgType": "share.file",
                "mediaUrls": ["https://file-stal-19.dlfl.vn/gr/abc"],
                "attachments": [{
                    "url": "https://file-stal-19.dlfl.vn/gr/abc",
                    "name": "22-KH.Tiếng nói xanh.pdf",
                    "mime": "application/pdf", "kind": "document",
                }],
            })

        self.assertEqual(len(handled), 1)
        event = handled[0]
        self.assertEqual(event.media_urls, ["C:/cache/documents/doc_22-KH.Tiếng nói xanh.pdf"])
        self.assertEqual(event.media_types, ["application/pdf"])
        self.assertEqual(event.message_type, zalo_adapter.MessageType.DOCUMENT)
        self.assertNotIn("ảnh", event.channel_context or "")
        # Người trong nhóm không có read_file, nên nội dung phải được kèm sẵn.
        self.assertIn("22-KH.Tiếng nói xanh.pdf", event.text)
        self.assertIn("KẾ HOẠCH tổ chức cuộc thi", event.text)
        self.assertEqual(event.media_text_inlined, [True])

    async def test_owner_dm_with_undownloadable_image_still_reaches_the_agent_with_the_reason(self):
        adapter = self.make_adapter()
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle

        async def fake_cache(url):
            raise ValueError("Inbound image payload is too large (99 bytes > 10 bytes)")

        with patch.object(adapter, "_is_owner", return_value=True), \
                patch.object(zalo_adapter, "cache_image_from_url", side_effect=fake_cache), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message(
                {
                    "type": "message",
                    "id": "dm-img",
                    "threadId": "1234567890123456789",
                    "threadType": zalo_adapter.THREAD_TYPE_USER,
                    "senderUid": "1234567890123456789",
                    "senderName": "Lương Hải Anh Cnt",
                    "text": "",
                    "msgType": "chat.photo",
                    "mediaUrls": ["https://example.com/big.jpg"],
                }
            )

        self.assertEqual(len(handled), 1)
        event = handled[0]
        self.assertEqual(event.media_urls, [])
        self.assertIn("[Người dùng gửi ảnh]", event.text)
        self.assertIn("quá dung lượng", event.channel_context)

    def test_image_failure_reason_distinguishes_format_network_and_size(self):
        reason = zalo_adapter._image_failure_reason
        non_image = ValueError("Refusing to cache non-image data as .jpg (starts with: 'x')")
        self.assertIn("JXL", reason("https://photo-stal-17.zdn.vn/gr/jxl/a/b", non_image))
        self.assertIn("không phải ảnh", reason("https://example.com/a", non_image))
        self.assertIn("quá dung lượng", reason("https://example.com/a", ValueError("Inbound image payload is too large (9 bytes > 1 bytes)")))
        self.assertIn("quá thời gian", reason("https://example.com/a", TimeoutError("timed out")))

        import httpx
        request = httpx.Request("GET", "https://example.com/a")
        gone = httpx.HTTPStatusError("404", request=request, response=httpx.Response(404, request=request))
        self.assertIn("HTTP 404", reason("https://example.com/a", gone))
        self.assertIn("quá thời gian", reason("https://example.com/a", httpx.ReadTimeout("slow", request=request)))

    async def test_inbound_dm_id_is_reused_as_dm_for_outbound_reply(self):
        adapter = self.make_adapter()
        adapter.handle_message = lambda _event: asyncio.sleep(0)

        sent = []

        async def fake_command(command, expect_ack=False):
            sent.append(command)
            return {"ok": True, "msgId": "reply-1"}

        adapter._command = fake_command
        with patch.object(adapter, "_is_owner", return_value=True), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message(
                {
                    "type": "message",
                    "id": "dm-1",
                    "threadId": "1234567890123456789",
                    "threadType": zalo_adapter.THREAD_TYPE_USER,
                    "senderUid": "1234567890123456789",
                    "senderName": "Lương Hải Anh Cnt",
                    "text": "Chào em",
                }
            )
            result = await adapter.send("1234567890123456789", "Em chào anh")

        self.assertTrue(result.success)
        self.assertEqual(sent[-1]["threadType"], zalo_adapter.THREAD_TYPE_USER)

    async def test_send_strips_hermes_plain_text_fallback_marker(self):
        adapter = self.make_adapter()
        sent = []

        async def fake_command(command, expect_ack=False):
            sent.append(command)
            return {"ok": True, "msgId": "reply-1"}

        adapter._command = fake_command
        result = await adapter.send(
            "1234567890123456789",
            "(Response formatting failed, plain text:)\n\nNội dung trả lời",
        )

        self.assertTrue(result.success)
        self.assertEqual(sent[-1]["text"], "Nội dung trả lời")

    async def test_send_strips_hermes_cron_wrapper_but_keeps_normal_text(self):
        adapter = self.make_adapter()
        sent = []

        async def fake_command(command, expect_ack=False):
            sent.append(command)
            return {"ok": True, "msgId": "cron-1"}

        adapter._command = fake_command
        wrapped = (
            "Cronjob Response: Nhắc họp\n(job_id: abc123)\n-------------\n\n"
            "Mai 7h30 họp chi đoàn nhé!\n\n"
            'To stop or manage this job, send me a new message (e.g. "stop reminder Nhắc họp").'
        )
        await adapter.send("9133571695356732407", wrapped, metadata={"chat_type": "group", "job_id": "abc123"})
        await adapter.send(
            "9133571695356732407", "Cronjob Response: là tên một mục trong báo cáo",
            metadata={"chat_type": "group"},
        )

        self.assertEqual(sent[0]["text"], "Mai 7h30 họp chi đoàn nhé!")
        self.assertEqual(sent[1]["text"], "Cronjob Response: là tên một mục trong báo cáo")

    async def test_send_drops_hermes_self_improvement_notice_but_keeps_normal_text(self):
        adapter = self.make_adapter()
        sent = []

        async def fake_command(command, expect_ack=False):
            sent.append(command)
            return {"ok": True, "msgId": "m1"}

        adapter._command = fake_command
        result = await adapter.send(
            "9133571695356732407", "💾 Self-improvement review: Skill 'zalo-chat-operations' patched",
            metadata={"chat_type": "group", "_interim_send": True},
        )
        self.assertTrue(result.success)
        self.assertEqual(sent, [])

        await adapter.send("9133571695356732407", "💾 là biểu tượng lưu tệp", metadata={"chat_type": "group"})
        self.assertEqual([c["text"] for c in sent], ["💾 là biểu tượng lưu tệp"])

    async def test_send_splits_new_message_marker_into_separate_messages(self):
        adapter = self.make_adapter()
        sent = []

        async def fake_command(command, expect_ack=False):
            sent.append(command)
            return {"ok": True, "msgId": f"m{len(sent)}"}

        adapter._command = fake_command
        result = await adapter.send(
            "9133571695356732407",
            "THÔNG BÁO\nMai họp chi đoàn lúc 7h30.\n\n[[NEW_MESSAGE]]\nEm soạn xong rồi ạ, anh xem tin trên nhé.",
            metadata={"chat_type": "group"},
        )

        self.assertTrue(result.success)
        self.assertEqual(
            [c["text"] for c in sent],
            ["THÔNG BÁO\nMai họp chi đoàn lúc 7h30.", "Em soạn xong rồi ạ, anh xem tin trên nhé."],
        )

    async def test_send_split_marker_never_leaks_when_the_model_formats_it_loosely(self):
        adapter = self.make_adapter()
        sent = []

        async def fake_command(command, expect_ack=False):
            sent.append(command)
            return {"ok": True, "msgId": f"m{len(sent)}"}

        adapter._command = fake_command
        for text in (
            "Bản soạn A\n**[[NEW_MESSAGE]]**\nXác nhận A",
            "Bản soạn B\r\n[[new_message]].\r\nXác nhận B",
            "Bản soạn C [[NEW MESSAGE]] Xác nhận C",
        ):
            await adapter.send("9133571695356732407", text, metadata={"chat_type": "group"})

        self.assertEqual(
            [c["text"] for c in sent],
            ["Bản soạn A", "Xác nhận A", "Bản soạn B", "Xác nhận B", "Bản soạn C", "Xác nhận C"],
        )

    async def test_owner_listed_in_ignore_sender_uids_is_still_answered(self):
        owner = "5736877140444221354"
        adapter = zalo_adapter.ZaloAdapter(
            PlatformConfig(
                enabled=True,
                extra={
                    "bridge_url": "ws://127.0.0.1:9",
                    "reply_only_tagged": True,
                    "ack_gestures": False,
                    "ignore_sender_uids": [owner],
                },
            )
        )
        adapter._self_profile = {"user_id": "bot-uid", "display_name": "Lăng Tiêu"}
        adapter._flood.check = lambda _uid: None
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle
        with patch.object(adapter, "_is_owner", side_effect=lambda uid: str(uid) == owner), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({
                "type": "message", "id": "g-own", "threadId": "g1", "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": owner, "senderName": "Chủ nhân", "text": "@Lăng Tiêu tóm tắt giúp",
                "mentions": [{"uid": "bot-uid"}],
            })
            await adapter._on_message({
                "type": "message", "id": "dm-own", "threadId": owner, "threadType": zalo_adapter.THREAD_TYPE_USER,
                "senderUid": owner, "senderName": "Chủ nhân", "text": "Chào em",
            })

        self.assertEqual([event.message_id for event in handled], ["g-own", "dm-own"])

    async def test_messages_from_ignored_bot_accounts_never_start_a_turn_but_stay_as_context(self):
        adapter = zalo_adapter.ZaloAdapter(
            PlatformConfig(
                enabled=True,
                extra={
                    "bridge_url": "ws://127.0.0.1:9",
                    "reply_only_tagged": True,
                    "ack_gestures": False,
                    "ignore_sender_uids": ["5736877140444221354"],
                },
            )
        )
        adapter._self_profile = {"user_id": "bot-uid", "display_name": "Lăng Tiêu"}
        adapter._flood.check = lambda _uid: None
        handled = []

        async def handle(event):
            handled.append(event)

        adapter.handle_message = handle
        frame = {
            "type": "message", "threadId": "g1", "threadType": zalo_adapter.THREAD_TYPE_GROUP,
            "text": "@Lăng Tiêu soi giúp ảnh này", "mentions": [{"uid": "bot-uid"}],
        }
        with patch.object(zalo_adapter, "_zalo_tools", return_value=DummyZaloTools()):
            await adapter._on_message({**frame, "id": "b1", "senderUid": "5736877140444221354", "senderName": "Uyển Nhi"})
            await adapter._on_message({**frame, "id": "m1", "senderUid": "3915070541883948642", "senderName": "Liên"})

        self.assertEqual([event.source.user_id for event in handled], ["3915070541883948642"])
        self.assertEqual(
            [entry["sender_uid"] for entry in adapter._recent_group_messages["g1"]],
            ["5736877140444221354", "3915070541883948642"],
        )

    async def test_turn_binding_failure_falls_back_to_public_not_system(self):
        adapter = self.make_adapter()

        class ExplodingTurns(dict):
            def get(self, *_args, **_kwargs):
                raise RuntimeError("hỏng bảng lượt")

        adapter._turns = ExplodingTurns()
        source = adapter.build_source(
            chat_id="group-1", chat_name="group-1", chat_type="group",
            user_id="2222222222222222222", user_name="M", message_id="m-x",
        )
        with patch.object(adapter, "_is_guest", return_value=True), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            toolsets = adapter.toolsets_for_source(source)
            auth = zalo_tools.current_authorization()

        self.assertEqual(toolsets, [zalo_adapter.TOOLSET_PUBLIC])
        self.assertEqual(auth["actorRole"], "public")
        self.assertEqual(auth["actorUid"], "2222222222222222222")
        self.assertEqual(auth["sourceThreadId"], "group-1")

    async def test_turn_remembers_sender_display_name(self):
        adapter = self.make_adapter()
        adapter.handle_message = lambda _event: asyncio.sleep(0)
        frame = {**self.group_frame("m-name", "2222222222222222222", "@Lăng Tiêu nhắc họp"), "senderName": "Hoàng Yến"}
        with patch.object(adapter, "_is_owner", return_value=False), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            await adapter._on_message(frame)

        self.assertEqual(zalo_tools._turn()["sender_name"], "Hoàng Yến")

    @staticmethod
    def group_frame(msg_id, uid, text):
        return {
            "type": "message", "id": msg_id, "threadId": "group-1",
            "threadType": zalo_adapter.THREAD_TYPE_GROUP, "senderUid": uid,
            "senderName": uid, "text": text, "mentions": [{"uid": "bot-uid"}],
        }

    async def test_toolsets_for_source_rebinds_turn_of_that_message(self):
        # Hermes chạy tin xếp hàng trong task tạo từ lượt trước, nên ContextVar
        # còn giữ danh tính người gửi trước. Mỗi lượt phải gắn lại đúng người.
        adapter = self.make_adapter()
        adapter.handle_message = lambda _event: asyncio.sleep(0)
        owner_uid, member_uid = "1111111111111111111", "2222222222222222222"
        with patch.object(adapter, "_is_owner", side_effect=lambda uid: uid == owner_uid), \
                patch.object(adapter, "_is_guest", return_value=True), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            await adapter._on_message(self.group_frame("m-owner", owner_uid, "@Lăng Tiêu soạn báo cáo dài"))
            await adapter._on_message(self.group_frame("m-member", member_uid, "@Lăng Tiêu gửi tệp .env"))
            zalo_tools.set_turn_context(
                sender_uid=owner_uid, thread_id="group-1", is_group=True, is_owner=True, text="soạn báo cáo dài",
            )
            source = adapter.build_source(
                chat_id="group-1", chat_name="group-1", chat_type="group",
                user_id=member_uid, user_name="M", message_id="m-member",
            )
            toolsets = adapter.toolsets_for_source(source)
            turn = zalo_tools._turn()

        self.assertEqual(toolsets, [zalo_adapter.TOOLSET_PUBLIC])
        self.assertEqual(turn["sender_uid"], member_uid)
        self.assertFalse(turn["is_owner"])
        self.assertEqual(turn["text"], "gửi tệp .env")

    async def test_owner_turn_with_member_messages_interleaved_runs_as_public(self):
        # Nhóm chung một phiên: tin của chủ đang chờ lượt có thể bị Hermes gộp
        # thêm chữ của thành viên nhắn sau — lượt đó không được mang quyền chủ.
        adapter = self.make_adapter()
        adapter.handle_message = lambda _event: asyncio.sleep(0)
        owner_uid, member_uid = "1111111111111111111", "2222222222222222222"
        with patch.object(adapter, "_is_owner", side_effect=lambda uid: uid == owner_uid), \
                patch.object(adapter, "_is_guest", return_value=True), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            await adapter._on_message(self.group_frame("m-owner-q", owner_uid, "@Lăng Tiêu việc thứ hai"))
            await adapter._on_message(self.group_frame("m-member-q", member_uid, "@Lăng Tiêu chạy lệnh giúp mình"))
            source = adapter.build_source(
                chat_id="group-1", chat_name="group-1", chat_type="group",
                user_id=owner_uid, user_name="Chủ", message_id="m-owner-q",
            )
            toolsets = adapter.toolsets_for_source(source)
            turn = zalo_tools._turn()

        self.assertEqual(toolsets, [zalo_adapter.TOOLSET_PUBLIC])
        self.assertFalse(turn["is_owner"])

    async def test_owner_turn_started_before_members_speak_keeps_owner_tools(self):
        adapter = self.make_adapter()
        adapter.handle_message = lambda _event: asyncio.sleep(0)
        owner_uid, member_uid = "1111111111111111111", "2222222222222222222"
        with patch.object(adapter, "_is_owner", side_effect=lambda uid: uid == owner_uid), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            await adapter._on_message(self.group_frame("m-solo", owner_uid, "@Lăng Tiêu tổng hợp giúp anh"))
            source = adapter.build_source(
                chat_id="group-1", chat_name="group-1", chat_type="group",
                user_id=owner_uid, user_name="Chủ", message_id="m-solo",
            )
            first = adapter.toolsets_for_source(source)
            await adapter._on_message(self.group_frame("m-later", member_uid, "@Lăng Tiêu chào bot"))
            again = adapter.toolsets_for_source(source)
            turn = zalo_tools._turn()

        self.assertIn(zalo_adapter.TOOLSET_OWNER, first)
        self.assertEqual(first, again)
        self.assertTrue(turn["is_owner"])

    async def test_toolsets_for_source_without_known_message_fails_closed(self):
        adapter = self.make_adapter()
        owner_uid = "1111111111111111111"
        with patch.object(adapter, "_is_owner", side_effect=lambda uid: uid == owner_uid), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            zalo_tools.set_turn_context(sender_uid="someone", thread_id="group-9", is_group=True, is_owner=True)
            source = adapter.build_source(
                chat_id="group-1", chat_name="group-1", chat_type="group",
                user_id=owner_uid, user_name="Chủ", message_id="unknown",
            )
            adapter.toolsets_for_source(source)
            turn = zalo_tools._turn()

        self.assertEqual(turn["sender_uid"], owner_uid)
        self.assertEqual(turn["thread_id"], "group-1")
        self.assertFalse(turn["is_owner"])

    async def test_group_turn_text_drops_bot_mention_so_confirmation_can_match(self):
        adapter = self.make_adapter()
        adapter.handle_message = lambda _event: asyncio.sleep(0)
        with patch.object(adapter, "_is_owner", return_value=True), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            await adapter._on_message(self.group_frame("m-confirm", "1111111111111111111", "@Lăng Tiêu  XÁC NHẬN A1B2C3"))

        self.assertEqual(zalo_tools._turn()["text"], "XÁC NHẬN A1B2C3")

    async def test_stranger_dm_sethome_gets_only_their_uid(self):
        adapter = self.make_adapter()
        handled = []

        async def fake_handle(event):
            handled.append(event)

        adapter.handle_message = fake_handle
        sent = []

        async def fake_command(command, expect_ack=False):
            sent.append((command, zalo_tools.current_authorization()))
            return {"ok": True, "msgId": "reply-1"}

        adapter._command = fake_command
        stranger = "3333333333333333333"
        with patch.object(adapter, "_is_owner", return_value=False), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            await adapter._on_message({
                "type": "message", "id": "dm-sethome", "threadId": stranger,
                "threadType": zalo_adapter.THREAD_TYPE_USER, "senderUid": stranger,
                "senderName": "Khách", "text": " /SetHome ",
            })

        self.assertEqual(handled, [])
        self.assertEqual(len(sent), 1)
        self.assertIn(stranger, sent[0][0]["text"])
        self.assertIn("chưa cấp quyền chủ", sent[0][0]["text"])
        self.assertEqual(sent[0][1]["actorUid"], stranger)
        self.assertEqual(sent[0][1]["actorRole"], "public")

    def test_owner_uid_is_a_dm_target_even_with_19_digits(self):
        adapter = self.make_adapter()
        owner_uid = "9000000000000000001"
        with patch.object(adapter, "_is_owner", side_effect=lambda uid: uid == owner_uid):
            self.assertEqual(adapter._guess_thread_type(owner_uid, {}), zalo_adapter.THREAD_TYPE_USER)
            self.assertEqual(adapter._guess_thread_type("9000000000000000002", {}), zalo_adapter.THREAD_TYPE_GROUP)

    def test_env_enablement_does_not_blank_out_config_yaml_values(self):
        # Hermes ghi kết quả _env_enablement ĐÈ lên extra của config.yaml. Trả
        # về bridge_token rỗng khi .env không đặt là xoá mất token trình cài ghi.
        keys = ("ZALO_BRIDGE_URL", "ZALO_BRIDGE_TOKEN", "ZALO_GROUP_REPLY_ONLY_TAGGED", "ZALO_HOME_CHANNEL")
        with patch.dict(os.environ, {}):
            for key in keys:
                os.environ.pop(key, None)
            seed = zalo_adapter._env_enablement() or {}
            self.assertNotIn("bridge_token", seed)
            self.assertNotIn("bridge_url", seed)
            self.assertNotIn("reply_only_tagged", seed)

            os.environ["ZALO_BRIDGE_TOKEN"] = "env-token"
            os.environ["ZALO_GROUP_REPLY_ONLY_TAGGED"] = "false"
            seed = zalo_adapter._env_enablement() or {}
            self.assertEqual(seed["bridge_token"], "env-token")
            self.assertIs(seed["reply_only_tagged"], False)

    def test_bridge_url_from_env_wins_over_config_extra(self):
        with patch.dict(os.environ, {"ZALO_BRIDGE_URL": "ws://127.0.0.1:3900"}):
            adapter = zalo_adapter.ZaloAdapter(
                PlatformConfig(enabled=True, extra={"bridge_url": "ws://127.0.0.1:3873"})
            )
        self.assertEqual(adapter._bridge_url, "ws://127.0.0.1:3900")

    async def test_send_voice_uploads_local_audio_then_forwards_zalo_cdn_url(self):
        # iPhone và Zalo PC không phát AAC thô qua link không đuôi: gửi M4A và nối đuôi.
        adapter = self.make_adapter()
        calls = []

        async def fake_invoke(method, args):
            calls.append((method, args))
            if method == "uploadAttachment":
                return {"ok": True, "result": [{"fileUrl": "https://fg41.dlfl.vn/abc/6538968052631542979"}]}
            return {"ok": True, "result": {"message": {"msgId": "voice-1"}}}

        adapter.invoke = fake_invoke
        with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as audio:
            audio.write(b"wav")
            audio_path = audio.name
        with tempfile.NamedTemporaryFile(suffix=".m4a", delete=False) as m4a:
            m4a.write(b"m4a")
            m4a_path = m4a.name
        try:
            with patch.object(zalo_adapter, "_transcode_to_m4a", return_value=m4a_path):
                result = await adapter.send_voice(
                    "2054797107487294899",
                    audio_path,
                    metadata={"chat_type": "group"},
                )
        finally:
            os.unlink(audio_path)

        self.assertTrue(result.success)
        self.assertEqual(calls[0][0], "uploadAttachment")
        self.assertTrue(calls[0][1][0][0].endswith(".m4a"))
        self.assertEqual(calls[0][1][1:], ["2054797107487294899", 1])
        self.assertEqual(calls[1], (
            "sendVoice",
            [
                {"voiceUrl": "https://fg41.dlfl.vn/abc/6538968052631542979.m4a", "ttl": 0},
                "2054797107487294899",
                1,
            ],
        ))
        self.assertFalse(os.path.exists(m4a_path), "tệp M4A tạm phải được dọn")

    async def test_send_voice_falls_back_to_aac_when_zalo_rejects_m4a(self):
        adapter = self.make_adapter()
        calls = []

        async def fake_invoke(method, args):
            calls.append((method, args))
            if method == "uploadAttachment":
                if args[0][0].endswith(".m4a"):
                    return {"ok": False, "error": 'File extension "m4a" is not allowed'}
                return {"ok": True, "result": [{"fileUrl": "https://fg41.dlfl.vn/abc/111"}]}
            return {"ok": True, "result": {"message": {"msgId": "voice-2"}}}

        adapter.invoke = fake_invoke
        with tempfile.NamedTemporaryFile(suffix=".aac", delete=False) as audio:
            audio.write(b"aac")
            audio_path = audio.name
        with tempfile.NamedTemporaryFile(suffix=".m4a", delete=False) as m4a:
            m4a.write(b"m4a")
            m4a_path = m4a.name
        try:
            with patch.object(zalo_adapter, "_transcode_to_m4a", return_value=m4a_path):
                result = await adapter.send_voice("2054797107487294899", audio_path,
                                                  metadata={"chat_type": "group"})
        finally:
            os.unlink(audio_path)

        self.assertTrue(result.success)
        self.assertEqual([c[0] for c in calls], ["uploadAttachment", "uploadAttachment", "sendVoice"])
        self.assertEqual(calls[2][1][0]["voiceUrl"], "https://fg41.dlfl.vn/abc/111.aac")

    async def test_send_voice_stops_on_network_error_instead_of_uploading_twice(self):
        adapter = self.make_adapter()
        calls = []

        async def fake_invoke(method, args):
            calls.append(method)
            return {"ok": False, "error": "timeout"}

        adapter.invoke = fake_invoke
        with tempfile.NamedTemporaryFile(suffix=".m4a", delete=False) as m4a:
            m4a.write(b"m4a")
            m4a_path = m4a.name
        with patch.object(zalo_adapter, "_transcode_to_m4a", return_value=m4a_path):
            result = await adapter.send_voice("2054797107487294899", m4a_path,
                                              metadata={"chat_type": "group"})
        self.assertFalse(result.success)
        self.assertEqual(calls, ["uploadAttachment"])

    async def test_send_voice_skips_same_file_resent_to_same_chat(self):
        # Bot gọi zalo_send_voice xong, gateway lại tự gắn MEDIA của text_to_speech vào câu trả
        # lời cuối và gửi lần nữa: cùng tệp, cùng nhóm thì chỉ được đi một lần.
        adapter = self.make_adapter()
        calls = []

        async def fake_invoke(method, args):
            calls.append((method, args[-2]))
            if method == "uploadAttachment":
                return {"ok": True, "result": [{"fileUrl": "https://fg41.dlfl.vn/abc/1"}]}
            return {"ok": True, "result": {"msgId": f"voice-{len(calls)}"}}

        def fake_m4a(_path):
            fd, path = tempfile.mkstemp(suffix=".m4a")
            os.close(fd)
            return path

        adapter.invoke = fake_invoke
        with tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as audio:
            audio.write(b"mp3")
            audio_path = audio.name
        try:
            with patch.object(zalo_adapter, "_transcode_to_m4a", side_effect=fake_m4a):
                first = await adapter.send_voice("6537986660262149071", audio_path,
                                                 metadata={"chat_type": "group"})
                again = await adapter.send_voice("6537986660262149071", audio_path,
                                                 metadata={"chat_type": "group"}, is_voice=True)
                other = await adapter.send_voice("39633293852382968", audio_path,
                                                 metadata={"chat_type": "group"})
        finally:
            os.unlink(audio_path)

        self.assertTrue(first.success and again.success and other.success)
        self.assertEqual(again.message_id, first.message_id)
        self.assertEqual(calls, [
            ("uploadAttachment", "6537986660262149071"), ("sendVoice", "6537986660262149071"),
            ("uploadAttachment", "39633293852382968"), ("sendVoice", "39633293852382968"),
        ])

    def test_with_audio_extension_only_adds_when_missing(self):
        self.assertEqual(zalo_adapter._with_audio_extension("https://fg41.dlfl.vn/a/1", ".m4a"),
                         "https://fg41.dlfl.vn/a/1.m4a")
        self.assertEqual(zalo_adapter._with_audio_extension("https://voice-aac-dl.zdn.vn/1/x.aac", ".m4a"),
                         "https://voice-aac-dl.zdn.vn/1/x.aac")

    def test_transcode_to_m4a_puts_moov_before_audio_data(self):
        import shutil
        import subprocess
        if not shutil.which("ffmpeg"):
            self.skipTest("máy này không có ffmpeg")
        fd, wav = tempfile.mkstemp(suffix=".wav")
        os.close(fd)
        try:
            subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=1",
                            wav], check=True, capture_output=True, timeout=60)
            out = zalo_adapter._transcode_to_m4a(wav)
            self.assertIsNotNone(out)
            data = open(out, "rb").read()
            os.unlink(out)
        finally:
            os.unlink(wav)
        self.assertEqual(data[4:8], b"ftyp")
        # +faststart: khối moov (thông tin tệp) phải đứng trước mdat (dữ liệu âm thanh).
        self.assertLess(data.find(b"moov"), data.find(b"mdat"))

    async def test_send_voice_accepts_media_delivery_is_voice_flag(self):
        adapter = self.make_adapter()

        async def fake_invoke(method, args):
            if method == "uploadAttachment":
                return {
                    "ok": True,
                    "result": [{"fileUrl": "https://cdn.zalo.test/voice.aac"}],
                }
            return {"ok": True, "result": {"msgId": "voice-media-1"}}

        adapter.invoke = fake_invoke
        with tempfile.NamedTemporaryFile(suffix=".aac", delete=False) as audio:
            audio.write(b"aac")
            audio_path = audio.name
        try:
            result = await adapter.send_voice(
                "9133571695356732407",
                audio_path,
                metadata={"chat_type": "group"},
                is_voice=True,
            )
        finally:
            os.unlink(audio_path)

        self.assertTrue(result.success)
        self.assertEqual(result.message_id, "voice-media-1")

    async def test_invoke_frame_carries_server_verifiable_turn_authorization(self):
        adapter = self.make_adapter()
        socket = CapturingSocket()
        adapter._ws = socket
        turn_token = zalo_tools._TURN.set({
            "sender_uid": "owner-1",
            "thread_id": "source-dm",
            "is_group": False,
            "is_owner": True,
            "text": "đổi tên nhóm",
        })
        try:
            pending = asyncio.create_task(adapter.invoke(
                "changeGroupName", ["Tên mới", "group-1"], confirmed=True,
            ))
            await asyncio.sleep(0)
            frame = socket.frames[0]
            await adapter._dispatch({"type": "ack", "reqId": frame["reqId"], "ok": True})
            await pending
        finally:
            zalo_tools._TURN.reset(turn_token)

        self.assertEqual(frame["auth"], {
            "actorUid": "owner-1",
            "actorRole": "owner",
            "sourceThreadId": "source-dm",
            "sourceThreadType": 0,
            "confirmed": True,
        })

    async def test_history_and_undo_frames_carry_owner_context_and_confirmation(self):
        adapter = self.make_adapter()
        socket = CapturingSocket()
        adapter._ws = socket
        turn_token = zalo_tools._TURN.set({
            "sender_uid": "owner-1", "thread_id": "group-1", "is_group": True,
            "is_owner": True, "text": "thu hồi tin vừa gửi",
        })
        try:
            history_task = asyncio.create_task(adapter.read_history(
                "group-1", 20, {"chat_type": "group"},
            ))
            await asyncio.sleep(0)
            history_frame = socket.frames[-1]
            await adapter._dispatch({"type": "ack", "reqId": history_frame["reqId"], "ok": True})
            await history_task

            undo_task = asyncio.create_task(adapter.undo_message(
                "group-1", metadata={"chat_type": "group"}, confirmed=True,
            ))
            await asyncio.sleep(0)
            undo_frame = socket.frames[-1]
            await adapter._dispatch({"type": "ack", "reqId": undo_frame["reqId"], "ok": True})
            await undo_task
        finally:
            zalo_tools._TURN.reset(turn_token)

        self.assertEqual(history_frame["auth"]["actorUid"], "owner-1")
        self.assertFalse(history_frame["auth"]["confirmed"])
        self.assertTrue(undo_frame["auth"]["confirmed"])

    async def test_application_heartbeat_sends_ping_without_authorization_payload(self):
        adapter = self.make_adapter()
        socket = CapturingSocket()
        adapter._ws = socket
        adapter._heartbeat_interval_s = 0.01
        task = asyncio.create_task(adapter._heartbeat_loop())
        try:
            await asyncio.sleep(0.025)
        finally:
            adapter._closing = True
            task.cancel()
            with self.assertRaises(asyncio.CancelledError):
                await task

        self.assertGreaterEqual(len(socket.frames), 1)
        self.assertEqual(socket.frames[0], {"type": "ping"})

    async def test_ack_gesture_uses_the_current_sender_authorization_context(self):
        adapter = self.make_adapter()
        adapter._ack_gestures = True
        adapter._auto_react = False
        adapter.handle_message = lambda _event: asyncio.sleep(0)
        observed = []

        async def fake_command(payload, expect_ack=False):
            observed.append((payload, zalo_tools.current_authorization()))
            return None

        adapter._command = fake_command
        with patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools), \
                patch.object(adapter, "_may_greet", return_value=True):
            await adapter._on_message({
                "type": "message", "id": "ack-1", "cliMsgId": "ack-c1",
                "threadId": "group-ack", "threadType": zalo_adapter.THREAD_TYPE_GROUP,
                "senderUid": "member-current", "senderName": "Thành viên",
                "text": "@Lăng Tiêu chào em", "mentions": [{"uid": "bot-uid"}],
            })

        self.assertEqual(observed[0][0]["type"], "ack_message")
        self.assertEqual(observed[0][1]["actorUid"], "member-current")
        self.assertEqual(observed[0][1]["sourceThreadId"], "group-ack")

    @requires_file_maker_deps
    def test_file_maker_builds_all_four_formats_with_vietnamese_text(self):
        def file_maker_black():
            from docx.shared import RGBColor
            return RGBColor(0, 0, 0)

        from plugins.zalo_tools import file_maker

        content = ("# Giáo án Hoá 10\n\n- **Mục tiêu:** hiểu H₂O\n  - ý con\n1. Khởi động\n\n"
                   "| Bước | Thời gian |\n|---|---|\n| Mở bài | 5 phút |\n\nĐoạn *nghiêng* cuối.")
        with tempfile.TemporaryDirectory() as tmp:
            made = {
                "docx": file_maker.make_file("docx", "Giáo án", directory=tmp, content=content),
                "pdf": file_maker.make_file("pdf", "Giáo án", directory=tmp, content=content),
                "pptx": file_maker.make_file("pptx", "Bài 7", directory=tmp, slides=[
                    {"title": "Sulfur", "bullets": ["Tính chất **vật lý**", "Ứng dụng"]}]),
                "xlsx": file_maker.make_file("xlsx", "Bảng điểm", directory=tmp, sheets=[
                    {"name": "Lớp 10A", "rows": [["Họ tên", "Điểm"], ["An", 9], ["Bình", "=1+1"]]}]),
            }
            for fmt, path in made.items():
                self.assertTrue(path.is_file() and path.stat().st_size > 0, fmt)
                self.assertEqual(path.parent, __import__("pathlib").Path(tmp))
                self.assertTrue(path.name.endswith("." + fmt))

            from docx import Document
            document = Document(str(made["docx"]))
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("Mục tiêu: hiểu H₂O", text)
            self.assertEqual(document.paragraphs[0].text, "Giáo án")
            section = document.sections[0]
            self.assertEqual(
                [round(m.cm, 1) for m in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin)],
                [2.0, 2.0, 3.0, 1.5], "lề theo Nghị định 30: 20/20/30/15 mm")
            body_xml = document.element.body.xml
            self.assertNotIn("w:numPr", body_xml, "ND30: không dùng danh sách tự động của Word")
            self.assertNotIn("w:shd", body_xml, "ND30: chữ và bảng đen trắng")
            self.assertIn("w:tblHeader", body_xml, "bảng lặp hàng tiêu đề mỗi trang")
            self.assertTrue(all(run.font.color.rgb in (None, file_maker_black())
                                for p in document.paragraphs for run in p.runs), "chữ màu đen")
            self.assertIn("- Mục tiêu: hiểu H₂O", text, "gạch đầu dòng gõ tay")
            from openpyxl import load_workbook
            sheet = load_workbook(str(made["xlsx"]))["Lớp 10A"]
            self.assertEqual(sheet["B3"].value, "'=1+1", "người lạ không được cài công thức Excel")
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(sheet.freeze_panes, "A2")
            from pptx import Presentation
            deck = Presentation(str(made["pptx"]))
            self.assertEqual(len(deck.slides), 2, "slide bìa + 1 slide nội dung")
            self.assertGreater(deck.slide_width, deck.slide_height, "khổ 16:9")

    def test_file_maker_rejects_bad_specs_and_unsafe_names(self):
        from plugins.zalo_tools import file_maker

        self.assertEqual(file_maker.safe_filename("../../etc/passwd", "pdf"), "etc passwd.pdf")
        self.assertEqual(file_maker.safe_filename("Giáo án: Hoá 10?", "docx"), "Giáo án Hoá 10.docx")
        self.assertEqual(file_maker.safe_filename("", "xlsx"), "Tài liệu.xlsx")
        with tempfile.TemporaryDirectory() as tmp:
            for fmt, kwargs in (
                ("exe", {"content": "x"}),
                ("docx", {"content": ""}),
                ("pdf", {"content": "a" * (file_maker.MAX_TEXT_CHARS + 1)}),
                ("pptx", {"slides": [{"title": "t", "bullets": ["b"] * 16}]}),
                ("xlsx", {"sheets": [{"name": "s", "rows": [["c"] * 31]}]}),
            ):
                with self.assertRaises(file_maker.FileSpecError, msg=fmt):
                    file_maker.make_file(fmt, "T", directory=tmp, **kwargs)

    @requires_file_maker_deps
    async def test_make_file_tool_sends_to_current_group_limits_members_and_cleans_up(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, args, confirmed=False):
                path = args[0]["attachments"][0]
                self.calls.append((method, args, os.path.isfile(path), path))
                return {"ok": True, "result": {"attachment": [{"msgId": "f1"}]}}

        fake = FakeAdapter()
        previous = zalo_tools._ACTIVE_ADAPTER
        zalo_tools._ACTIVE_ADAPTER = fake
        zalo_tools._FILE_QUOTA.clear()
        member = {"sender_uid": "9000000000000000001", "thread_id": "7903718250465581275",
                  "is_group": True, "is_owner": False, "text": "tạo giúp file"}
        spec = {"thread_id": "7903718250465581275", "format": "docx", "title": "Đề kiểm tra",
                "content": "# Câu 1\n- Ý a"}
        try:
            token = zalo_tools._TURN.set(member)
            try:
                results = [json.loads(await zalo_tools.zalo_make_file(dict(spec))) for _ in range(6)]
                other_group = json.loads(await zalo_tools.zalo_make_file(dict(spec, thread_id="1111111111111111111")))
            finally:
                zalo_tools._TURN.reset(token)
            token = zalo_tools._TURN.set(dict(member, is_group=False, thread_id="9000000000000000001"))
            try:
                dm = json.loads(await zalo_tools.zalo_make_file(dict(spec, thread_id="9000000000000000001")))
            finally:
                zalo_tools._TURN.reset(token)
        finally:
            zalo_tools._ACTIVE_ADAPTER = previous
            zalo_tools._FILE_QUOTA.clear()

        self.assertTrue(all(r["success"] for r in results[:5]))
        self.assertFalse(results[5]["success"])
        self.assertIn("5 tệp mỗi giờ", results[5]["error"])
        self.assertFalse(other_group["success"])
        self.assertFalse(dm["success"])
        self.assertEqual(len(fake.calls), 5)
        method, args, existed, path = fake.calls[0]
        self.assertEqual(method, "sendMessage")
        self.assertEqual(args[1:], ["7903718250465581275", 1])
        self.assertTrue(existed and path.endswith("Đề kiểm tra.docx"))
        self.assertFalse(os.path.exists(path), "thư mục tạm phải được xoá sau khi gửi")

    async def test_poll_vote_and_add_option_tools_send_numeric_ids(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, args, confirmed=False):
                self.calls.append((method, args))
                return {"ok": True, "result": {"options": []}}

        fake = FakeAdapter()
        previous = zalo_tools._ACTIVE_ADAPTER
        zalo_tools._ACTIVE_ADAPTER = fake
        try:
            voted = json.loads(await zalo_tools.zalo_vote_poll(
                {"poll_id": "1137063889", "option_ids": ["1137063892"]}))
            withdrawn = json.loads(await zalo_tools.zalo_vote_poll({"poll_id": "1137063889", "option_ids": []}))
            added = json.loads(await zalo_tools.zalo_add_poll_options(
                {"poll_id": "1137063889", "options": [" Tôi là bot ", ""], "vote": True}))
            bad = json.loads(await zalo_tools.zalo_vote_poll({"poll_id": "1137063889", "option_ids": ["abc"]}))
            no_option = json.loads(await zalo_tools.zalo_add_poll_options({"poll_id": "1137063889", "options": []}))
        finally:
            zalo_tools._ACTIVE_ADAPTER = previous

        self.assertTrue(voted["success"] and withdrawn["success"] and added["success"])
        self.assertFalse(bad["success"])
        self.assertFalse(no_option["success"])
        self.assertEqual(fake.calls, [
            ("votePoll", [1137063889, [1137063892]]),
            ("votePoll", [1137063889, []]),
            ("addPollOptions", [{"pollId": 1137063889,
                                 "options": [{"voted": True, "content": "Tôi là bot"}],
                                 "votedOptionIds": []}]),
        ])

    async def test_zalo_send_voice_accepts_tts_local_path(self):
        class FakeAdapter:
            async def send_voice(self, chat_id, audio_path, metadata=None):
                self.call = (chat_id, audio_path, metadata)
                return zalo_adapter.SendResult(success=True, message_id="voice-2")

            async def invoke(self, method, args):
                return {"ok": False, "error": "local paths are not public URLs"}

        fake = FakeAdapter()
        with tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as audio:
            audio.write(b"mp3")
            audio_path = audio.name
        previous = zalo_tools._ACTIVE_ADAPTER
        zalo_tools._ACTIVE_ADAPTER = fake
        turn_token = zalo_tools._TURN.set({
            "sender_uid": "1234567890123456789",
            "thread_id": "1234567890123456789",
            "is_group": False,
            "is_owner": True,
            "text": "Gửi voice vào nhóm Đệ ruột",
        })
        try:
            response = await zalo_tools.zalo_send_voice({
                "thread_id": 2054797107487294899,
                "thread_kind": "group",
                "url": audio_path,
            })
        finally:
            zalo_tools._TURN.reset(turn_token)
            zalo_tools._ACTIVE_ADAPTER = previous
            os.unlink(audio_path)

        self.assertTrue(json.loads(response)["success"])
        self.assertEqual(fake.call, (
            "2054797107487294899",
            audio_path,
            {"chat_type": "group"},
        ))

    async def test_ack_wakes_a_command_waiting_on_another_event_loop(self):
        # Công cụ chạy trên vòng lặp riêng của luồng agent, ack tới trên vòng lặp
        # gateway. Ack phải đánh thức lệnh ngay, không để nó chờ hết thời gian chờ.
        adapter = self.make_adapter()
        socket = CapturingSocket()
        adapter._ws = socket
        result = {}

        def tool_thread():
            loop = asyncio.new_event_loop()
            try:
                started = time.monotonic()
                result["ack"] = loop.run_until_complete(adapter._command(
                    {"type": "history", "threadId": "g1", "threadType": 1}, expect_ack=True,
                ))
                result["seconds"] = time.monotonic() - started
            finally:
                loop.close()

        with patch.object(zalo_adapter, "ACK_TIMEOUT_SECONDS", 3), \
                patch.object(zalo_adapter, "_zalo_tools", return_value=zalo_tools):
            thread = threading.Thread(target=tool_thread)
            thread.start()
            for _ in range(200):
                if socket.frames:
                    break
                await asyncio.sleep(0.01)
            await adapter._dispatch({"type": "ack", "reqId": socket.frames[0]["reqId"], "ok": True})
            await asyncio.to_thread(thread.join)

        self.assertEqual(result["ack"], {"type": "ack", "reqId": socket.frames[0]["reqId"], "ok": True})
        self.assertLess(result["seconds"], 1.0)


    async def test_platform_tools_enforce_owner_file_guest_and_stranger_tiers(self):
        from hermes_cli.tools_config import (_get_platform_tools,
                                             _get_plugin_toolset_keys)
        from toolsets import resolve_toolset

        # Bài này đo tầng quyền qua registry thật, nên nó cần plugin Zalo đã
        # được khám phá. Đo ngày 18/09/2026: _get_plugin_toolset_keys() đọc bộ
        # khoá mà lần chạy trước đã lưu trong HOME của Hermes
        # (tools_config.py:158, get_plugin_toolset_keys_nowait). Trong một
        # container trống với HOME sạch, bộ đó rỗng → ``zalo_public`` giải ra 0
        # công cụ và bài kiểm thử đỏ vì thiếu môi trường, chứ không vì quyền sai.
        # Chạy trong container hermes đang phục vụ (env thật) thì nó giải đúng
        # 16. Bỏ qua kèm lý do thay vì hạ con số 16 xuống cho xanh.
        if zalo_tools.TOOLSET_PUBLIC not in _get_plugin_toolset_keys():
            self.skipTest(
                "plugin Zalo chưa được khám phá trong môi trường này — "
                "chạy bài này trong container hermes có HOME thật"
            )

        zalo_tools.define_platform_composite()
        zalo_tools.define_denied_toolset()
        adapter = self.make_adapter()
        owner_uid = "9000000000000000001"
        guest_uid = "9000000000000000002"
        stranger_uid = "9000000000000000003"

        def effective_tools(uid):
            source = adapter.build_source(
                chat_id="9000000000000000004", chat_name="group-1", chat_type="group",
                user_id=uid, user_name=uid, message_id=f"message-{uid}",
            )
            override = adapter.toolsets_for_source(source)
            # ``known_plugin_toolsets`` phải có. Thiếu khoá này thì
            # _get_platform_tools trả về cả ``zalo_owner`` cho lượt khách,
            # biến fixture thành một thế giới lỏng hơn runtime. Config đã triển
            # khai có khoá này; với nó, khách đo đúng 16 công cụ và không có
            # ``terminal``. Đừng tăng count để che mất ranh giới đó.
            probe = {
                "platform_toolsets": {"zalo": override},
                "known_plugin_toolsets": {
                    "zalo": [zalo_tools.TOOLSET_OWNER, zalo_tools.TOOLSET_PUBLIC,
                             zalo_tools.TOOLSET_CRON],
                },
            }
            return {tool for toolset in _get_platform_tools(probe, "zalo")
                    for tool in resolve_toolset(toolset)}

        with tempfile.TemporaryDirectory() as directory:
            roster_path = os.path.join(directory, "roster.json")
            with open(roster_path, "w", encoding="utf-8") as roster_file:
                json.dump({
                    "version": 1,
                    "owners": [owner_uid],
                    "guests": [guest_uid],
                }, roster_file)
            with open(os.path.join(directory, "guest-groups.json"), "w", encoding="utf-8") as groups_file:
                json.dump({"version": 1, "guestGroups": ["9000000000000000004"]}, groups_file)

            with patch.dict(os.environ, {
                "ZALO_ALLOWED_USERS": owner_uid,
                "GATEWAY_ALLOWED_USERS": "",
                "ZALO_ROSTER_FILE": roster_path,
                "ZALO_ALLOW_ALL_USERS": "true",
            }), patch.object(adapter, "_bind_turn_for_source", side_effect=lambda _source, uid: uid == owner_uid):
                owner_tools = effective_tools(owner_uid)
                guest_tools = effective_tools(guest_uid)
                stranger_tools = effective_tools(stranger_uid)
                self.assertFalse(set(zalo_tools.ZALO_DENIED_CORE_TOOLS) & owner_tools)
                self.assertEqual(len(guest_tools), 16)
                self.assertFalse(set(zalo_tools.ZALO_DENIED_CORE_TOOLS) & guest_tools)
                self.assertEqual(stranger_tools, set())
                self.assertTrue(adapter._may_greet(guest_uid))
                self.assertFalse(adapter._may_greet(stranger_uid))


class ZaloToolSchemaTest(unittest.TestCase):
    def test_every_zalo_tool_has_one_unique_public_or_owner_assignment(self):
        names = [name for name, _emoji, _schema, _handler, _toolset in zalo_tools.TOOLS]
        assignments = [toolset for _name, _emoji, _schema, _handler, toolset in zalo_tools.TOOLS]
        self.assertEqual(len(names), len(set(names)))
        self.assertEqual(set(assignments), {
            zalo_tools.TOOLSET_PUBLIC, zalo_tools.TOOLSET_OWNER, zalo_tools.TOOLSET_CRON,
        })
        self.assertEqual(assignments.count(zalo_tools.TOOLSET_PUBLIC), 16)
        # Owner tool count excludes guest-user lifecycle tools. Any new owner
        # tool must update this explicit boundary assertion.
        self.assertEqual(assignments.count(zalo_tools.TOOLSET_OWNER), 34)
        self.assertEqual(assignments.count(zalo_tools.TOOLSET_CRON), 1)

    def test_zalo_ids_remain_strings_through_hermes_argument_coercion(self):
        import model_tools

        original = "2054797107487294899"
        schema = next(
            schema for name, _emoji, schema, _handler, _toolset in zalo_tools.TOOLS
            if name == "zalo_undo"
        )
        with patch.object(model_tools.registry, "get_schema", return_value=schema):
            coerced = model_tools.coerce_tool_args("zalo_undo", {
                "thread_id": original,
                "thread_kind": "group",
            })

        self.assertEqual(coerced["thread_id"], original)
        self.assertIsInstance(coerced["thread_id"], str)

    def test_send_voice_requires_string_zalo_thread_id(self):
        schema = next(
            schema for name, _emoji, schema, _handler, _toolset in zalo_tools.TOOLS
            if name == "zalo_send_voice"
        )["parameters"]

        errors = list(Draft7Validator(schema).iter_errors({
            "thread_id": "2054797107487294899",
            "thread_kind": "group",
            "url": "https://example.com/voice.aac",
        }))

        self.assertEqual(errors, [])

    def test_all_zalo_id_fields_are_declared_as_strings(self):
        for name, _emoji, schema, _handler, _toolset in zalo_tools.TOOLS:
            properties = schema["parameters"].get("properties", {})
            for key, spec in properties.items():
                if not (key.endswith("_id") or key.endswith("_ids")):
                    continue
                item_spec = spec.get("items") if spec.get("type") == "array" else spec
                with self.subTest(tool=name, field=key):
                    self.assertEqual(item_spec.get("type"), "string")

    def test_dangerous_owner_tools_expose_human_confirmation_code(self):
        dangerous = {
            "zalo_lock_poll", "zalo_pin_conversation", "zalo_mute", "zalo_undo",
            "zalo_rename_group", "zalo_group_member_change", "zalo_group_deputy",
            "zalo_review_member", "zalo_create_group", "zalo_invite_to_groups",
            "zalo_group_link", "zalo_join_group_link", "zalo_set_bio",
            "zalo_set_active_status",
        }
        schemas = {name: schema["parameters"] for name, _emoji, schema, _handler, _toolset in zalo_tools.TOOLS}
        for name in dangerous:
            with self.subTest(tool=name):
                self.assertNotIn("confirmation_code", schemas[name]["required"])
                self.assertEqual(schemas[name]["properties"]["confirmation_code"]["type"], "string")

    def test_cron_member_toolset_holds_only_safe_group_tools(self):
        from toolsets import resolve_toolset

        zalo_tools.define_cron_member_toolset()

        self.assertEqual(set(resolve_toolset(zalo_tools.TOOLSET_CRON_MEMBER, include_registry=False)), {
            "zalo_web_search", "zalo_web_read", "zalo_kb_list", "zalo_kb_read", "zalo_group_history",
        })

    def test_no_mcp_sentinel_in_per_job_toolsets_adds_no_mcp_servers(self):
        from cron.scheduler import _resolve_cron_enabled_toolsets

        self.assertEqual(
            _resolve_cron_enabled_toolsets({"enabled_toolsets": ["zalo_cron_member", "no_mcp"]}, {}),
            ["zalo_cron_member"],
        )


class ZaloToolContractTest(unittest.IsolatedAsyncioTestCase):
    async def test_owner_tool_fails_closed_without_turn_context(self):
        raw_handler = next(
            handler for name, _emoji, _schema, handler, _toolset in zalo_tools.TOOLS
            if name == "zalo_list_people"
        )
        handler = zalo_tools._owner_only(raw_handler, "zalo_list_people")
        token = zalo_tools._TURN.set({})
        try:
            response = await handler({})
        finally:
            zalo_tools._TURN.reset(token)
        self.assertFalse(json.loads(response)["success"])
        self.assertIn("chủ nhân", json.loads(response)["error"])

    async def test_public_voice_cannot_read_local_file_outside_knowledge_base(self):
        class FakeAdapter:
            async def send_voice(self, *_args, **_kwargs):
                raise AssertionError("unsafe local file reached adapter")

        with tempfile.TemporaryDirectory() as knowledge_base, \
                tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as audio:
            audio.write(b"mp3")
            audio_path = audio.name
        previous = zalo_tools._ACTIVE_ADAPTER
        zalo_tools._ACTIVE_ADAPTER = FakeAdapter()
        turn_token = zalo_tools._TURN.set({
            "sender_uid": "public-user", "thread_id": "group-1",
            "is_group": True, "is_owner": False, "text": "gửi voice",
        })
        try:
            with patch.dict(os.environ, {"ZALO_KB_DIR": knowledge_base}):
                response = await zalo_tools.zalo_send_voice({
                    "thread_id": "group-1", "thread_kind": "group", "url": audio_path,
                })
        finally:
            zalo_tools._TURN.reset(turn_token)
            zalo_tools._ACTIVE_ADAPTER = previous
            os.unlink(audio_path)
        self.assertFalse(json.loads(response)["success"])
        self.assertIn("kho tài liệu", json.loads(response)["error"])

    async def test_public_voice_rejects_private_network_url(self):
        class FakeAdapter:
            async def invoke(self, *_args, **_kwargs):
                raise AssertionError("private voice URL reached the sidecar")

        zalo_tools._ACTIVE_ADAPTER = FakeAdapter()
        token = zalo_tools._TURN.set({
            "sender_uid": "public-user", "thread_id": "group-1",
            "is_group": True, "is_owner": False, "text": "gửi voice",
        })
        try:
            response = await zalo_tools.zalo_send_voice({
                "thread_id": "group-1", "thread_kind": "group", "url": "http://127.0.0.1:8080/secret.aac",
            })
        finally:
            zalo_tools._TURN.reset(token)
        self.assertFalse(json.loads(response)["success"])

    async def test_group_members_asks_bridge_for_that_group(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def group_members(self, chat_id):
                self.calls.append(chat_id)
                return {"ok": True, "result": {"total": 1, "members": [{"id": "u1", "displayName": "An"}]}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        token = zalo_tools._TURN.set({
            "sender_uid": "public-user", "thread_id": "group-1",
            "is_group": True, "is_owner": False, "text": "nhóm có ai",
        })
        try:
            response = await zalo_tools.zalo_group_members({"thread_id": "group-1"})
        finally:
            zalo_tools._TURN.reset(token)
        self.assertTrue(json.loads(response)["success"], response)
        self.assertEqual(fake.calls, ["group-1"])

    async def test_fb_publish_can_schedule_a_post(self):
        from plugins.zalo_tools import facebook as zalo_fb

        calls = []

        def fake_graph(path, method="GET", **params):
            params.pop("timeout", None)
            calls.append((path, method, params))
            return {"id": "page-1_post-1"} if method == "POST" else {"permalink_url": "https://fb.test/p"}

        draft = {"page": {"id": "page-1", "token": "page-token", "name": "Trang thử"}, "message": "Nội dung", "photos": []}
        with patch.object(zalo_fb, "confirmed_in_message", return_value=True), \
                patch.object(zalo_fb, "take_draft", return_value=(draft, None)), \
                patch.object(zalo_fb, "graph", side_effect=fake_graph):
            response = await zalo_tools.zalo_fb_publish({"code": "ABC123", "scheduled_publish_time": "1790000000"})

        result = json.loads(response)
        self.assertTrue(result["success"], response)
        self.assertTrue(result["result"]["len_lich"])
        self.assertEqual(calls[0][1], "POST")
        self.assertEqual(calls[0][2]["published"], "false")
        self.assertEqual(calls[0][2]["scheduled_publish_time"], 1790000000)

    async def test_fb_publish_rejects_bad_schedule_before_consuming_draft(self):
        from plugins.zalo_tools import facebook as zalo_fb

        with patch.object(zalo_fb, "confirmed_in_message", return_value=True), \
                patch.object(zalo_fb, "take_draft", side_effect=AssertionError("draft consumed")):
            response = await zalo_tools.zalo_fb_publish({"code": "ABC123", "scheduled_publish_time": "ngày mai"})

        self.assertFalse(json.loads(response)["success"])

    def test_authorization_outside_a_chat_turn_is_system(self):
        token = zalo_tools._TURN.set(None)
        try:
            auth = zalo_tools.current_authorization()
        finally:
            zalo_tools._TURN.reset(token)
        self.assertEqual(auth["actorRole"], "system")
        self.assertEqual(auth["actorUid"], "")

    def setUp(self):
        self.previous_adapter = zalo_tools._ACTIVE_ADAPTER
        self.turn_token = zalo_tools._TURN.set({
            "sender_uid": "1234567890123456789",
            "thread_id": "1234567890123456789",
            "is_group": False,
            "is_owner": True,
            "text": "test",
        })

    def tearDown(self):
        zalo_tools._TURN.reset(self.turn_token)
        zalo_tools._ACTIVE_ADAPTER = self.previous_adapter

    async def test_sticker_maps_search_result_to_send_payload(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, args):
                self.calls.append((method, args))
                if method == "searchSticker":
                    return {"ok": True, "result": [
                        {"sticker_id": "123", "cate_id": "45", "type": "3"},
                    ]}
                return {"ok": True, "result": {"msgId": 999}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        response = await zalo_tools.zalo_send_sticker({
            "thread_id": 9133571695356732407,
            "thread_kind": "group",
            "keyword": "ngủ ngon",
        })

        self.assertTrue(json.loads(response)["success"])
        self.assertEqual(fake.calls[-1], (
            "sendSticker",
            [{"id": 123, "cateId": 45, "type": 3}, "9133571695356732407", 1],
        ))

    async def test_send_file_and_link_normalize_numeric_thread_id(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, args):
                self.calls.append((method, args))
                return {"ok": True, "result": {"message": {"msgId": 999}}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as item:
            item.write(b"test")
            item_path = item.name
        try:
            file_response = await zalo_tools.zalo_send_file({
                "thread_id": 9133571695356732407,
                "thread_kind": "group",
                "path": item_path,
            })
        finally:
            os.unlink(item_path)
        link_response = await zalo_tools.zalo_send_link({
            "thread_id": 9133571695356732407,
            "thread_kind": "group",
            "url": "https://example.com",
        })

        self.assertTrue(json.loads(file_response)["success"])
        self.assertTrue(json.loads(link_response)["success"])
        self.assertEqual(fake.calls[0][1][1], "9133571695356732407")
        self.assertEqual(fake.calls[1][1][1], "9133571695356732407")

    async def test_group_member_change_normalizes_numeric_ids(self):
        class FakeAdapter:
            async def invoke(self, method, args):
                self.call = (method, args)
                return {"ok": True, "result": {"status": 0}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        response = await zalo_tools.zalo_group_member_change({
            "group_id": 9133571695356732407,
            "user_ids": [1234567890123456789],
            "action": "add",
        })

        self.assertTrue(json.loads(response)["success"])
        self.assertEqual(fake.call, (
            "addUserToGroup",
            [["1234567890123456789"], "9133571695356732407"],
        ))

    async def test_other_group_management_tools_normalize_numeric_ids(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, args):
                self.calls.append((method, args))
                return {"ok": True, "result": {"status": 0}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        await zalo_tools.zalo_group_deputy({
            "group_id": 9133571695356732407,
            "user_id": 1234567890123456789,
            "action": "add",
        })
        await zalo_tools.zalo_review_member({
            "group_id": 9133571695356732407,
            "user_ids": [1234567890123456789],
            "approve": True,
        })
        await zalo_tools.zalo_create_group({"member_ids": [1234567890123456789]})
        await zalo_tools.zalo_invite_to_groups({
            "user_id": 1234567890123456789,
            "group_ids": [9133571695356732407],
        })

        self.assertEqual(fake.calls, [
            ("addGroupDeputy", ["1234567890123456789", "9133571695356732407"]),
            ("reviewPendingMemberRequest", [{
                "members": ["1234567890123456789"], "isApprove": True,
            }, "9133571695356732407"]),
            ("createGroup", [{"members": ["1234567890123456789"]}]),
            ("inviteUserToGroups", [
                "1234567890123456789", ["9133571695356732407"],
            ]),
        ])

    async def test_read_history_uses_adapter_history_contract(self):
        class FakeAdapter:
            async def read_history(self, chat_id, count, metadata=None):
                self.call = (chat_id, count, metadata)
                return {"ok": True, "result": {"messages": [{"msgId": "m1"}]}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        response = await zalo_tools.zalo_read_history({
            "thread_id": 9133571695356732407,
            "thread_kind": "group",
            "count": 20,
        })

        self.assertTrue(json.loads(response)["success"])
        self.assertEqual(fake.call, (
            "9133571695356732407", 20, {"chat_type": "group"},
        ))

    async def test_read_history_since_hours_pages_until_done_and_formats_lines(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def read_history_range(self, chat_id, since_ms, until_ms, cursor=None, limit=300, metadata=None):
                self.calls.append((chat_id, cursor, limit, metadata))
                pages = {
                    None: ({"messages": [
                        {"ts": 1789290000000, "senderName": "Yến", "text": "mai họp mấy giờ", "isSelf": False},
                        {"ts": 1789290060000, "senderName": "", "senderUid": "u9", "text": "dòng 1\ndòng 2"},
                    ]}, "1789290060000:2"),
                    "1789290060000:2": ({"messages": [
                        {"ts": 1789290120000, "text": "8h nhé", "isSelf": True},
                        {"ts": 1789290180000, "senderName": "Trang", "text": "", "msgType": "chat.sticker"},
                    ]}, None),
                }
                result, next_cursor = pages[cursor]
                return {"ok": True, "result": {**result, "nextCursor": next_cursor}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        response = json.loads(await zalo_tools.zalo_read_history({
            "thread_id": 39633293852382968, "thread_kind": "group", "since_hours": 24,
        }))

        self.assertTrue(response["success"])
        data = response["data"] if "data" in response else response
        payload = data.get("result", data)
        self.assertEqual(payload["count"], 4)
        self.assertFalse(payload["con_nua"])
        self.assertEqual([c[1] for c in fake.calls], [None, "1789290060000:2"])
        self.assertEqual(fake.calls[0][0], "39633293852382968")
        lines = payload["text"].split("\n")
        self.assertRegex(lines[0], r"^\[\d\d/\d\d \d\d:\d\d\] Yến: mai họp mấy giờ$")
        self.assertTrue(lines[1].endswith("u9: dòng 1 / dòng 2"))
        self.assertTrue(lines[2].endswith("Bot: 8h nhé"))
        self.assertTrue(lines[3].endswith("Trang: [chat.sticker]"))

    async def test_read_history_range_stops_at_char_budget_and_hands_back_a_cursor(self):
        class FakeAdapter:
            async def read_history_range(self, chat_id, since_ms, until_ms, cursor=None, limit=300, metadata=None):
                n = int(cursor or 0)
                return {"ok": True, "result": {
                    "messages": [{"ts": 1789290000000 + i, "senderName": "A", "text": "x" * 900} for i in range(limit)],
                    "nextCursor": str(n + 1),
                }}

        zalo_tools._ACTIVE_ADAPTER = FakeAdapter()
        response = json.loads(await zalo_tools.zalo_read_history({
            "thread_id": 39633293852382968, "thread_kind": "group", "since_hours": 24,
        }))
        payload = response.get("data", response)
        payload = payload.get("result", payload)
        self.assertTrue(payload["con_nua"])
        self.assertTrue(payload["next_cursor"])
        self.assertIn("cursor", payload["huong_dan"])
        self.assertLess(len(payload["text"]), zalo_tools.HISTORY_RANGE_CHAR_BUDGET + 300 * 1100)

    async def test_undo_without_ids_uses_latest_own_message_contract(self):
        class FakeAdapter:
            async def undo_message(self, chat_id, msg_id=None, cli_msg_id=None, metadata=None, *, confirmed=False):
                self.call = (chat_id, msg_id, cli_msg_id, metadata, confirmed)
                return {"ok": True, "result": {"status": 0, "msgId": "m1"}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        response = await zalo_tools.zalo_undo({
            "thread_id": 9133571695356732407,
            "thread_kind": "group",
        })

        self.assertTrue(json.loads(response)["success"])
        self.assertEqual(fake.call, (
            "9133571695356732407", None, None, {"chat_type": "group"}, False,
        ))

    async def test_undo_without_ids_targets_the_quoted_own_message(self):
        class FakeAdapter:
            async def undo_message(self, chat_id, msg_id=None, cli_msg_id=None, metadata=None, *, confirmed=False):
                self.call = (chat_id, msg_id, cli_msg_id, metadata, confirmed)
                return {"ok": True, "result": {"status": 0, "msgId": msg_id}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        quoted_turn = zalo_tools._TURN.set({
            "sender_uid": "owner", "thread_id": "2054797107487294899",
            "is_group": True, "is_owner": True, "text": "thu hồi tin nhắn này",
            "reply_msg_id": "8240551224624", "reply_cli_msg_id": "1788864027075",
            "reply_is_own": True,
        })
        try:
            response = await zalo_tools.zalo_undo({
                "thread_id": "2054797107487294899", "thread_kind": "group",
            })
        finally:
            zalo_tools._TURN.reset(quoted_turn)

        self.assertTrue(json.loads(response)["success"])
        self.assertEqual(fake.call, (
            "2054797107487294899", "8240551224624", "1788864027075",
            {"chat_type": "group"}, False,
        ))

    async def test_owner_dangerous_action_runs_immediately_by_default(self):
        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, args, *, confirmed=False):
                self.calls.append((method, args, confirmed))
                return {"ok": True, "result": {"status": 0}}

        self.enterContext(patch.dict(os.environ, {}))
        os.environ.pop("ZALO_CONFIRM_DANGEROUS", None)
        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        guarded = zalo_tools._confirmed_action(
            zalo_tools.zalo_group_member_change, "zalo_group_member_change",
        )
        args = {"group_id": "g1", "user_ids": ["u1"], "action": "remove"}

        owner_turn = zalo_tools._TURN.set({
            "sender_uid": "owner", "thread_id": "g1", "is_group": True,
            "is_owner": True, "text": "xoá u1 khỏi nhóm",
        })
        try:
            owner = json.loads(await guarded(args))
        finally:
            zalo_tools._TURN.reset(owner_turn)

        member_turn = zalo_tools._TURN.set({
            "sender_uid": "member", "thread_id": "g1", "is_group": True,
            "is_owner": False, "text": "xoá u1 khỏi nhóm",
        })
        try:
            member = json.loads(await guarded(args))
        finally:
            zalo_tools._TURN.reset(member_turn)

        self.assertTrue(owner["success"], owner)
        self.assertFalse(member["success"])
        self.assertEqual(fake.calls, [("removeUserFromGroup", [["u1"], "g1"], True)])

    async def test_undo_confirmation_keeps_the_quoted_target_on_the_later_turn(self):
        self.enterContext(patch.dict(os.environ, {"ZALO_CONFIRM_DANGEROUS": "true"}))

        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def undo_message(self, chat_id, msg_id=None, cli_msg_id=None, metadata=None, *, confirmed=False):
                self.calls.append((chat_id, msg_id, cli_msg_id, metadata, confirmed))
                return {"ok": True, "result": {"status": 0, "msgId": msg_id}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        guarded = zalo_tools._confirmed_action(zalo_tools.zalo_undo, "zalo_undo")
        args = {"thread_id": "2054797107487294899", "thread_kind": "group"}
        first_turn = zalo_tools._TURN.set({
            "sender_uid": "owner-quoted", "thread_id": "2054797107487294899",
            "is_group": True, "is_owner": True, "text": "thu hồi tin nhắn này",
            "reply_msg_id": "8240551224624", "reply_cli_msg_id": "1788864027075",
            "reply_is_own": True,
        })
        try:
            challenge = json.loads(await guarded(args))
        finally:
            zalo_tools._TURN.reset(first_turn)

        second_turn = zalo_tools._TURN.set({
            "sender_uid": "owner-quoted", "thread_id": "2054797107487294899",
            "is_group": True, "is_owner": True,
            "text": f'XÁC NHẬN {challenge["confirmation_code"]}',
        })
        try:
            result = json.loads(await guarded({
                **args, "confirmation_code": challenge["confirmation_code"],
            }))
        finally:
            zalo_tools._TURN.reset(second_turn)

        self.assertTrue(result["success"])
        self.assertEqual(fake.calls, [(
            "2054797107487294899", "8240551224624", "1788864027075",
            {"chat_type": "group"}, True,
        )])

    async def test_confirmation_guard_requires_code_in_a_later_owner_message(self):
        self.enterContext(patch.dict(os.environ, {"ZALO_CONFIRM_DANGEROUS": "true"}))

        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, args, *, confirmed=False):
                self.calls.append((method, args, confirmed))
                return {"ok": True, "result": {"status": 0}}

        fake = FakeAdapter()
        zalo_tools._ACTIVE_ADAPTER = fake
        guarded = zalo_tools._confirmed_action(
            zalo_tools.zalo_group_member_change, "zalo_group_member_change",
        )
        args = {"group_id": "g1", "user_ids": ["u1"], "action": "remove"}
        first_turn = zalo_tools._TURN.set({
            "sender_uid": "owner", "thread_id": "dm-owner", "is_group": False,
            "is_owner": True, "text": "xóa u1 khỏi nhóm g1",
        })
        try:
            challenge = json.loads(await guarded(args))
            same_turn = json.loads(await guarded({**args, "confirmation_code": challenge["confirmation_code"]}))
        finally:
            zalo_tools._TURN.reset(first_turn)
        second_turn = zalo_tools._TURN.set({
            "sender_uid": "owner", "thread_id": "dm-owner", "is_group": False,
            "is_owner": True, "text": f'XÁC NHẬN {challenge["confirmation_code"]}',
        })
        try:
            allowed = json.loads(await guarded({**args, "confirmation_code": challenge["confirmation_code"]}))
        finally:
            zalo_tools._TURN.reset(second_turn)

        self.assertFalse(challenge["success"])
        self.assertFalse(same_turn["success"])
        self.assertTrue(allowed["success"])
        self.assertEqual(fake.calls, [
            ("removeUserFromGroup", [["u1"], "g1"], True),
        ])

    async def test_confirmation_code_rejects_negation_and_changed_arguments(self):
        self.enterContext(patch.dict(os.environ, {"ZALO_CONFIRM_DANGEROUS": "true"}))

        class FakeAdapter:
            def __init__(self):
                self.calls = []

            async def invoke(self, method, params, *, confirmed=False):
                self.calls.append((method, params, confirmed))
                return {"ok": True, "result": {"status": 0}}

        fake = FakeAdapter()
        previous_adapter = zalo_tools._ACTIVE_ADAPTER
        zalo_tools._ACTIVE_ADAPTER = fake
        guarded = zalo_tools._confirmed_action(
            zalo_tools.zalo_group_member_change, "zalo_group_member_change",
        )
        args = {"group_id": "g1", "user_ids": ["u1"], "action": "remove"}
        first_turn = zalo_tools._TURN.set({
            "sender_uid": "owner-2", "thread_id": "dm-owner-2", "is_group": False,
            "is_owner": True, "text": "xóa u1 khỏi nhóm g1",
        })
        try:
            challenge = json.loads(await guarded(args))
        finally:
            zalo_tools._TURN.reset(first_turn)
        code = challenge["confirmation_code"]

        negated_turn = zalo_tools._TURN.set({
            "sender_uid": "owner-2", "thread_id": "dm-owner-2", "is_group": False,
            "is_owner": True, "text": f"KHÔNG XÁC NHẬN {code}",
        })
        try:
            negated = json.loads(await guarded({**args, "confirmation_code": code}))
        finally:
            zalo_tools._TURN.reset(negated_turn)

        changed_turn = zalo_tools._TURN.set({
            "sender_uid": "owner-2", "thread_id": "dm-owner-2", "is_group": False,
            "is_owner": True, "text": f"XÁC NHẬN {code}",
        })
        try:
            changed = json.loads(await guarded({
                **args, "group_id": "g2", "confirmation_code": code,
            }))
        finally:
            zalo_tools._TURN.reset(changed_turn)
            zalo_tools._ACTIVE_ADAPTER = previous_adapter

        self.assertFalse(negated["success"])
        self.assertFalse(changed["success"])
        self.assertEqual(fake.calls, [])


class ZaloCronTurnTest(unittest.IsolatedAsyncioTestCase):
    OWNER = "9200000000000000001"
    GROUP = "9133000000000000001"
    MEMBER = "3900000000000000001"

    def setUp(self):
        self.turn_token = zalo_tools._TURN.set(None)
        self.previous_adapter = zalo_tools._ACTIVE_ADAPTER
        zalo_tools._ACTIVE_ADAPTER = None
        self.enterContext(patch.dict(os.environ, {"ZALO_ALLOWED_USERS": self.OWNER}))

    def tearDown(self):
        zalo_tools._TURN.reset(self.turn_token)
        zalo_tools._ACTIVE_ADAPTER = self.previous_adapter

    def fake_jobs(self):
        return FakeCronJobs([
            {"id": "owner-job", "deliver": f"zalo:{self.GROUP}",
             "origin": {"platform": "zalo", "chat_id": self.GROUP, "user_id": "someone-else"}},
            {"id": "group-job", "deliver": f"zalo:{self.GROUP}",
             "origin": {"platform": "zalo", "chat_id": self.GROUP, "chat_type": "group",
                        "zalo_scope": "group", "zalo_creator_uid": self.MEMBER, "zalo_creator_name": "Yến"}},
            {"id": "telegram-job", "deliver": "telegram:8617174143",
             "origin": {"platform": "telegram", "chat_id": "8617174143"}},
        ])

    @staticmethod
    async def probe(_args, **_kw):
        return json.dumps({"turn": zalo_tools._turn(), "auth": zalo_tools.current_authorization()})

    async def call(self, task_id, jobs):
        guarded = zalo_tools._with_cron_turn(self.probe, "probe")
        with patch.object(zalo_tools, "_cron_jobs", return_value=jobs):
            return json.loads(await guarded({}, task_id=task_id))

    async def test_owner_created_cron_runs_as_owner_in_its_target_chat(self):
        seen = await self.call("cron:owner-job:run-1", self.fake_jobs())

        self.assertTrue(seen["turn"]["is_owner"])
        self.assertEqual(seen["turn"]["sender_uid"], self.OWNER)
        self.assertEqual(seen["turn"]["thread_id"], self.GROUP)
        self.assertTrue(seen["turn"]["is_group"])
        self.assertEqual(seen["turn"]["text"], "")
        self.assertEqual(seen["auth"]["actorRole"], "owner")
        self.assertEqual(seen["auth"]["sourceThreadType"], zalo_tools.THREAD_GROUP)
        self.assertEqual(seen["auth"]["cronJobId"], "owner-job")

    async def test_group_cron_runs_as_its_creator_locked_to_that_group(self):
        seen = await self.call("cron:group-job:run-1", self.fake_jobs())

        self.assertFalse(seen["turn"]["is_owner"])
        self.assertEqual(seen["turn"]["sender_uid"], self.MEMBER)
        self.assertEqual(seen["turn"]["sender_name"], "Yến")
        self.assertEqual(seen["auth"]["actorRole"], "public")
        self.assertEqual(seen["auth"]["sourceThreadId"], self.GROUP)
        self.assertEqual(seen["auth"]["cronJobId"], "group-job")

    async def test_corrupted_group_marker_never_becomes_an_owner_turn(self):
        jobs = FakeCronJobs([
            {"id": "scope-typo", "deliver": f"zalo:{self.GROUP}",
             "origin": {"platform": "zalo", "chat_id": self.GROUP, "zalo_scope": "Group",
                        "zalo_creator_uid": self.MEMBER}},
            {"id": "creator-only", "deliver": f"zalo:{self.GROUP}",
             "origin": {"platform": "zalo", "chat_id": self.GROUP, "zalo_creator_uid": self.MEMBER}},
        ])
        for task_id in ("cron:scope-typo:run-1", "cron:creator-only:run-1"):
            with self.subTest(task_id=task_id):
                seen = await self.call(task_id, jobs)
                self.assertEqual(seen["turn"], {})
                self.assertEqual(seen["auth"]["actorRole"], "system")

    async def test_non_cron_task_missing_job_or_non_zalo_job_get_no_turn(self):
        jobs = self.fake_jobs()
        for task_id in ("session-abc", "cron:missing:run-1", "cron:telegram-job:run-1", ""):
            with self.subTest(task_id=task_id):
                seen = await self.call(task_id, jobs)
                self.assertEqual(seen["turn"], {})
                self.assertEqual(seen["auth"]["actorRole"], "system")

    async def test_owner_cron_without_allowlist_gets_no_owner_turn(self):
        with patch.dict(os.environ, {"ZALO_ALLOWED_USERS": ""}):
            seen = await self.call("cron:owner-job:run-1", self.fake_jobs())

        self.assertEqual(seen["turn"], {})

    async def test_existing_chat_turn_is_never_replaced(self):
        chat = zalo_tools._TURN.set({
            "sender_uid": self.MEMBER, "thread_id": "other", "is_group": True, "is_owner": False, "text": "hi",
        })
        try:
            seen = await self.call("cron:owner-job:run-1", self.fake_jobs())
        finally:
            zalo_tools._TURN.reset(chat)

        self.assertFalse(seen["turn"]["is_owner"])
        self.assertEqual(seen["turn"]["thread_id"], "other")

    async def test_turn_is_restored_after_the_cron_call(self):
        await self.call("cron:owner-job:run-1", self.fake_jobs())

        self.assertEqual(zalo_tools._turn(), {})

    async def test_registered_owner_tool_works_inside_owner_cron(self):
        class FakeAdapter:
            async def read_history(self, chat_id, count, metadata=None):
                self.call = (chat_id, count, metadata)
                return {"ok": True, "result": {"messages": []}}

        ctx, fake = FakeToolContext(), FakeAdapter()
        zalo_tools.register_tools(ctx)
        zalo_tools._ACTIVE_ADAPTER = fake
        with patch.object(zalo_tools, "_cron_jobs", return_value=self.fake_jobs()):
            response = await ctx.handlers["zalo_read_history"](
                {"thread_id": self.GROUP, "thread_kind": "group", "count": 5},
                task_id="cron:owner-job:run-1",
            )

        self.assertTrue(json.loads(response)["success"], response)
        self.assertEqual(fake.call, (self.GROUP, 5, {"chat_type": "group"}))

    async def test_group_history_reads_only_the_cron_group_and_ignores_model_thread_id(self):
        class FakeAdapter:
            async def read_history(self, chat_id, count, metadata=None):
                self.call = (chat_id, count, metadata)
                return {"ok": True, "result": {"messages": [{"msgId": "m1"}]}}

        ctx, fake = FakeToolContext(), FakeAdapter()
        zalo_tools.register_tools(ctx)
        zalo_tools._ACTIVE_ADAPTER = fake
        with patch.object(zalo_tools, "_cron_jobs", return_value=self.fake_jobs()):
            response = await ctx.handlers["zalo_group_history"](
                {"count": 500, "thread_id": "other-group"}, task_id="cron:group-job:run-1",
            )

        self.assertTrue(json.loads(response)["success"], response)
        self.assertEqual(fake.call, (self.GROUP, 100, {"chat_type": "group"}))

    async def test_group_history_refuses_outside_cron(self):
        chat = zalo_tools._TURN.set({
            "sender_uid": self.MEMBER, "thread_id": self.GROUP, "is_group": True, "is_owner": False, "text": "đọc nhóm",
        })
        try:
            response = await zalo_tools.zalo_group_history({})
        finally:
            zalo_tools._TURN.reset(chat)

        self.assertFalse(json.loads(response)["success"])

    async def test_allowed_owner_uids_fails_closed_when_multiplexed_and_unscoped(self):
        import agent.secret_scope

        with patch.object(
            agent.secret_scope, "get_secret",
            side_effect=agent.secret_scope.UnscopedSecretError("no scope"),
        ):
            self.assertEqual(zalo_tools._allowed_owner_uids(), [])


class ZaloGroupCronTest(unittest.IsolatedAsyncioTestCase):
    OWNER = "9200000000000000001"
    GROUP = "9133000000000000001"
    OTHER_GROUP = "9133000000000000002"
    MEMBER = "3900000000000000001"
    OTHER_MEMBER = "3900000000000000002"

    def setUp(self):
        self.turn_token = zalo_tools._TURN.set(None)
        self.enterContext(patch.dict(os.environ, {"ZALO_ALLOWED_USERS": self.OWNER}))

    def tearDown(self):
        zalo_tools._TURN.reset(self.turn_token)

    def member_turn(self, uid=None, *, is_owner=False, is_group=True, group=None, **extra):
        return {
            "sender_uid": uid or self.MEMBER, "sender_name": "Yến", "thread_id": group or self.GROUP,
            "is_group": is_group, "is_owner": is_owner, "text": "hẹn giờ", **extra,
        }

    async def run_tool(self, args, jobs, turn):
        token = zalo_tools._TURN.set(turn)
        try:
            with patch.object(zalo_tools, "_cron_jobs", return_value=jobs):
                return json.loads(await zalo_tools.zalo_group_cron(args))
        finally:
            zalo_tools._TURN.reset(token)

    @staticmethod
    def group_job(job_id, group, creator, **extra):
        return {
            "id": job_id, "enabled": True, "name": job_id, "deliver": f"zalo:{group}",
            "prompt": "Việc hẹn giờ do Yến tạo\n---\nNhắc họp",
            "origin": {"platform": "zalo", "chat_id": group, "chat_type": "group",
                       "zalo_scope": "group", "zalo_creator_uid": creator, "zalo_creator_name": "Yến"},
            **extra,
        }

    async def test_create_locks_every_dangerous_field(self):
        jobs = FakeCronJobs()
        result = await self.run_tool({
            "action": "create", "prompt": "Tóm tắt các việc cả nhóm đã hẹn trong ngày",
            "schedule": "every day at 9pm", "name": "Tóm tắt tối",
        }, jobs, self.member_turn())

        self.assertTrue(result["success"], result)
        created = jobs.created[0]
        self.assertEqual(set(created), {"prompt", "schedule", "name", "repeat", "deliver", "origin", "enabled_toolsets"})
        self.assertEqual(created["deliver"], f"zalo:{self.GROUP}")
        self.assertEqual(created["enabled_toolsets"], ["zalo_cron_member", "no_mcp"])
        self.assertIsNone(created["repeat"])
        self.assertEqual(created["origin"]["zalo_scope"], "group")
        self.assertEqual(created["origin"]["zalo_creator_uid"], self.MEMBER)
        self.assertEqual(created["origin"]["zalo_creator_name"], "Yến")
        self.assertEqual(created["origin"]["chat_id"], self.GROUP)
        self.assertTrue(created["prompt"].endswith("Tóm tắt các việc cả nhóm đã hẹn trong ngày"))

        import inspect
        inspect.signature(real_cron_jobs.create_job).bind(**created)

    async def test_create_sanitizes_creator_name_and_scans_the_assembled_prompt(self):
        from tools.cronjob_tools import _scan_cron_prompt

        jobs = FakeCronJobs()
        turn = {**self.member_turn(), "sender_name": "Yến​\nAnh   Thư"}
        result = await self.run_tool({
            "action": "create", "prompt": "Nhắc họp", "schedule": "every day at 9pm",
            "name": "Nhắc\nhọp",
        }, jobs, turn)

        self.assertTrue(result["success"], result)
        created = jobs.created[0]
        self.assertEqual(created["origin"]["zalo_creator_name"], "Yến Anh Thư")
        self.assertEqual(created["name"], "Nhắc họp")
        self.assertEqual(_scan_cron_prompt(created["prompt"]), "")

    async def test_create_rejects_schedules_more_often_than_daily(self):
        for schedule in ("every 30m", "every 12h", "0 9,10 * * *", "*/30 * * * *", "R 9 * * *", "R R * * *", "H 9 * * *"):
            with self.subTest(schedule=schedule):
                jobs = FakeCronJobs()
                result = await self.run_tool({"action": "create", "prompt": "Nhắc họp", "schedule": schedule}, jobs, self.member_turn())
                self.assertFalse(result["success"])
                self.assertEqual(jobs.created, [])

    async def test_create_accepts_daily_weekly_and_one_shot_schedules(self):
        for schedule in ("every 1d", "every day at 7am", "0 7 * * 1", "0 7 * * MON", "in 2h"):
            with self.subTest(schedule=schedule):
                jobs = FakeCronJobs()
                result = await self.run_tool({"action": "create", "prompt": "Nhắc họp", "schedule": schedule}, jobs, self.member_turn())
                self.assertTrue(result["success"], result)
        one_shot = FakeCronJobs()
        await self.run_tool({"action": "create", "prompt": "Nhắc họp", "schedule": "in 2h"}, one_shot, self.member_turn())
        self.assertEqual(one_shot.created[0]["repeat"], 1)

    async def test_create_rejects_past_one_shot_long_prompt_and_bad_schedule(self):
        cases = (
            {"prompt": "Nhắc họp", "schedule": "2020-01-01T07:30"},
            {"prompt": "x" * 1001, "schedule": "every day at 7am"},
            {"prompt": "Nhắc họp", "schedule": "hôm nào đó"},
            {"prompt": "", "schedule": "every day at 7am"},
        )
        for case in cases:
            with self.subTest(case=case["schedule"]):
                jobs = FakeCronJobs()
                result = await self.run_tool({"action": "create", **case}, jobs, self.member_turn())
                self.assertFalse(result["success"])
                self.assertEqual(jobs.created, [])

    async def test_create_enforces_quota_per_member_and_per_group_but_not_for_owner(self):
        mine = FakeCronJobs([self.group_job(f"m{i}", self.OTHER_GROUP, self.MEMBER) for i in range(3)])
        result = await self.run_tool({"action": "create", "prompt": "Nhắc họp", "schedule": "every day at 7am"}, mine, self.member_turn())
        self.assertFalse(result["success"])
        self.assertIn("3", result["error"])

        crowded = [self.group_job(f"g{i}", self.GROUP, f"39000000000000001{i:02d}") for i in range(10)]
        result = await self.run_tool({"action": "create", "prompt": "Nhắc họp", "schedule": "every day at 7am"}, FakeCronJobs(crowded), self.member_turn())
        self.assertFalse(result["success"])
        self.assertIn("10", result["error"])

        finished = FakeCronJobs([
            self.group_job("m0", self.OTHER_GROUP, self.MEMBER),
            self.group_job("m1", self.OTHER_GROUP, self.MEMBER),
            self.group_job("m2", self.OTHER_GROUP, self.MEMBER, state="completed"),
        ])
        result = await self.run_tool({"action": "create", "prompt": "Nhắc họp", "schedule": "every day at 7am"}, finished, self.member_turn())
        self.assertTrue(result["success"], result)

        owner = await self.run_tool(
            {"action": "create", "prompt": "Nhắc họp", "schedule": "every day at 7am"},
            FakeCronJobs(crowded), self.member_turn(self.OWNER, is_owner=True),
        )
        self.assertTrue(owner["success"], owner)

    async def test_tool_works_only_in_groups_and_never_inside_cron(self):
        args = {"action": "create", "prompt": "Nhắc họp", "schedule": "every day at 7am"}
        dm = await self.run_tool(args, FakeCronJobs(), self.member_turn(is_group=False))
        in_cron = await self.run_tool(args, FakeCronJobs(), self.member_turn(cron_job_id="group-job"))
        listing_in_cron = await self.run_tool({"action": "list"}, FakeCronJobs(), self.member_turn(cron_job_id="group-job"))

        self.assertFalse(dm["success"])
        self.assertFalse(in_cron["success"])
        self.assertFalse(listing_in_cron["success"])

    async def test_create_uses_hermes_prompt_scanner_and_refuses_when_it_is_missing(self):
        args = {"action": "create", "prompt": "Nhắc họp", "schedule": "every day at 7am"}
        with patch("tools.cronjob_tools._scan_cron_prompt", return_value="Blocked: threat"):
            blocked = await self.run_tool(args, FakeCronJobs(), self.member_turn())
        with patch.dict(sys.modules, {"tools.cronjob_tools": None}):
            missing = await self.run_tool(args, FakeCronJobs(), self.member_turn())

        self.assertFalse(blocked["success"])
        self.assertIn("Blocked", blocked["error"])
        self.assertFalse(missing["success"])

    async def test_list_shows_this_group_and_hides_owner_prompts(self):
        jobs = FakeCronJobs([
            {"id": "owner-job", "enabled": True, "name": "Bản tin", "deliver": f"zalo:{self.GROUP}",
             "prompt": "bí mật của chủ nhân", "origin": {"platform": "zalo", "chat_id": self.GROUP}},
            self.group_job("group-job", self.GROUP, self.MEMBER),
            self.group_job("elsewhere", self.OTHER_GROUP, self.MEMBER),
        ])
        result = await self.run_tool({"action": "list"}, jobs, self.member_turn(self.OTHER_MEMBER))

        self.assertTrue(result["success"], result)
        items = {item["job_id"]: item for item in result["result"]["jobs"]}
        self.assertEqual(set(items), {"owner-job", "group-job"})
        self.assertEqual(items["owner-job"]["nguoi_tao"], "chủ nhân")
        self.assertNotIn("noi_dung", items["owner-job"])
        self.assertEqual(items["group-job"]["noi_dung"], "Nhắc họp")
        self.assertNotIn("bí mật", json.dumps(result, ensure_ascii=False))

    async def test_remove_follows_creator_or_owner_rule(self):
        def jobs():
            return FakeCronJobs([
                {"id": "owner-job", "enabled": True, "deliver": f"zalo:{self.GROUP}",
                 "origin": {"platform": "zalo", "chat_id": self.GROUP}},
                self.group_job("group-job", self.GROUP, self.MEMBER),
                self.group_job("elsewhere", self.OTHER_GROUP, self.MEMBER),
            ])

        other = jobs()
        self.assertFalse((await self.run_tool({"action": "remove", "job_id": "group-job"}, other, self.member_turn(self.OTHER_MEMBER)))["success"])
        self.assertEqual(other.removed, [])

        creator = jobs()
        self.assertTrue((await self.run_tool({"action": "remove", "job_id": "group-job"}, creator, self.member_turn()))["success"])
        self.assertEqual(creator.removed, ["group-job"])

        member_vs_owner = jobs()
        self.assertFalse((await self.run_tool({"action": "remove", "job_id": "owner-job"}, member_vs_owner, self.member_turn()))["success"])
        self.assertEqual(member_vs_owner.removed, [])

        owner = jobs()
        self.assertTrue((await self.run_tool({"action": "remove", "job_id": "owner-job"}, owner, self.member_turn(self.OWNER, is_owner=True)))["success"])
        self.assertEqual(owner.removed, ["owner-job"])

        wrong_group = jobs()
        self.assertFalse((await self.run_tool({"action": "remove", "job_id": "elsewhere"}, wrong_group, self.member_turn()))["success"])
        self.assertEqual(wrong_group.removed, [])


class ZaloMemberToolGuardTest(unittest.TestCase):
    """Hermes ghim bộ công cụ của phiên nhóm theo lượt đầu (thường là chủ nhân)
    rồi cấp lại cho mọi lượt sau, bất kể toolsets_for_source trả gì. Hook
    pre_tool_call là rào chắn tại điểm thực thi."""

    MEMBER = "3900000000000000001"
    GROUP = "9133000000000000001"

    def setUp(self):
        self.turn_token = zalo_tools._TURN.set(None)

    def tearDown(self):
        zalo_tools._TURN.reset(self.turn_token)

    def guard(self, tool_name, args=None):
        return zalo_tools.guard_member_tool_call(
            tool_name=tool_name, args=args if args is not None else {}, task_id="t", session_id="s", tool_call_id="c",
        )

    def bind_member(self):
        zalo_tools.bind_turn({"sender_uid": self.MEMBER, "thread_id": self.GROUP,
                              "is_group": True, "is_owner": False, "text": ""})

    def test_every_zalo_role_denies_generic_mutation_tools(self):
        for turn in (
            {"sender_uid": self.MEMBER, "thread_id": self.GROUP, "is_group": True, "is_owner": False, "text": ""},
            {"sender_uid": "9200000000000000001", "thread_id": self.GROUP, "is_group": True, "is_owner": True, "text": ""},
        ):
            zalo_tools.bind_turn(turn)
            for name in zalo_tools.ZALO_DENIED_CORE_TOOLS:
                verdict = self.guard(name)
                self.assertEqual(verdict["action"], "block", name)
                self.assertEqual(verdict["message"], "Hành động này không khả dụng qua Zalo.")

    def test_member_turn_keeps_public_zalo_mcp_and_tool_search_bridge(self):
        self.bind_member()
        public = next(name for name, _e, _s, _h, ts in zalo_tools.TOOLS if ts == zalo_tools.TOOLSET_PUBLIC)
        owner_only = next(name for name, _e, _s, _h, ts in zalo_tools.TOOLS if ts == zalo_tools.TOOLSET_OWNER)
        self.assertIsNone(self.guard(public))
        self.assertIsNone(self.guard("tool_search"))
        self.assertEqual(self.guard(owner_only)["action"], "block")

        from tools.registry import registry
        with patch.object(registry, "get_toolset_for_tool", return_value="mcp-rag"):
            self.assertIsNone(self.guard("rag_search"))

    def test_non_zalo_context_is_untouched_and_owner_keeps_narrow_group_lifecycle(self):
        self.assertIsNone(self.guard("terminal"))
        zalo_tools.bind_turn(None)
        self.assertIsNone(self.guard("terminal"))
        zalo_tools.bind_turn({"sender_uid": "9200000000000000001", "thread_id": self.GROUP,
                              "is_group": True, "is_owner": True, "text": ""})
        self.assertIsNone(self.guard("zalo_grant_guest_group"))
        self.assertIsNone(self.guard("zalo_revoke_guest_group"))

    def test_plugin_entry_registers_the_guard_as_pre_tool_call_hook(self):
        import plugins.zalo_tools as plugin

        ctx = FakeToolContext()
        with patch.object(plugin, "define_platform_composite"), \
                patch.object(plugin, "define_cron_member_toolset"):
            plugin.register(ctx)
        self.assertEqual(ctx.hooks.get("pre_tool_call"), [zalo_tools.guard_member_tool_call])

    def test_owner_turn_blocks_generic_core_even_without_outsider(self):
        class FakeAdapter:
            pass

        adapter = FakeAdapter()
        adapter._turns = {"m1": {"thread_id": self.GROUP, "is_owner": True, "seq": 5}}
        public = next(name for name, _e, _s, _h, ts in zalo_tools.TOOLS if ts == zalo_tools.TOOLSET_PUBLIC)
        with patch.object(zalo_tools, "_ACTIVE_ADAPTER", adapter), \
                patch.dict(os.environ, {"HERMES_GATEWAY_BUSY_INPUT_MODE": "queue"}):
            zalo_tools.bind_turn({"sender_uid": "9200000000000000001", "thread_id": self.GROUP,
                                  "is_group": True, "is_owner": True, "text": "", "seq": 5})
            self.assertEqual(self.guard("terminal")["action"], "block")
            self.assertIsNone(self.guard(public))

    def test_wrapped_generic_core_action_is_denied_for_every_zalo_role(self):
        self.bind_member()
        verdict = self.guard(
            "tool_call", {"name": "terminal", "arguments": {"command": "cat .env"}},
        )
        self.assertEqual(verdict["action"], "block")

        zalo_tools.bind_turn({"sender_uid": "9200000000000000001", "thread_id": self.GROUP,
                              "is_group": True, "is_owner": True, "text": ""})
        verdict = self.guard(
            "tool_call", {"name": "skill_manage", "arguments": {"action": "write"}},
        )
        self.assertEqual(verdict["action"], "block")

    def test_unmatched_owner_message_never_restores_generic_core_tools(self):
        adapter = ZaloAdapterMediaContextTest.make_adapter(self)
        bound = []

        class CapturingTools:
            def bind_turn(self, turn):
                bound.append(turn)

        class Source:
            def __init__(self, chat_type):
                self.user_id = "9200000000000000001"
                self.chat_id = "chat-1"
                self.chat_type = chat_type
                self.message_id = "not-a-received-message"

        with patch.object(zalo_adapter, "_zalo_tools", return_value=CapturingTools()), \
                patch.object(adapter, "_is_owner", return_value=True):
            dm = adapter.toolsets_for_source(Source("dm"))
            group = adapter.toolsets_for_source(Source("group"))

        self.assertEqual(dm, [zalo_tools.TOOLSET_OWNER, zalo_tools.TOOLSET_PUBLIC])
        self.assertEqual(group, [zalo_tools.TOOLSET_DENIED])
        zalo_tools.bind_turn(bound[0])
        self.assertEqual(self.guard("terminal")["action"], "block")
        zalo_tools.bind_turn(bound[1])
        self.assertEqual(self.guard("terminal")["action"], "block")


class FacebookVisibilityTest(unittest.TestCase):
    """Graph API vẫn khai 'đã đăng, công khai' cho một bài mà người ngoài không
    xem được — đúng chuyện đã xảy ra ngày 09/09. Chỉ trình nhúng công khai mới
    nói thật, nên kiểm tra bằng nó trước khi báo chủ nhân là đã đăng."""

    def visibility(self, html, status=200):
        from plugins.zalo_tools import facebook as fb

        class FakeResponse:
            def __enter__(self_inner):
                return self_inner

            def __exit__(self_inner, *_a):
                return False

            def read(self_inner):
                return html.encode("utf-8")

        if status != 200:
            with patch.object(fb.urllib.request, "urlopen", side_effect=OSError("mạng hỏng")):
                return fb.public_visibility("https://facebook.com/1/posts/2")
        with patch.object(fb.urllib.request, "urlopen", return_value=FakeResponse()):
            return fb.public_visibility("https://facebook.com/1/posts/2")

    def test_embed_says_post_is_gone(self):
        html = ("<div>Bài viết này trên Facebook không còn nữa vì có thể đã bị gỡ hoặc "
                "cài đặt quyền riêng tư của bài viết đã thay đổi.</div>")
        self.assertEqual(self.visibility(html), "khong_xem_duoc")
        self.assertEqual(self.visibility("<p>This content isn't available right now</p>"), "khong_xem_duoc")

    def test_rendered_embed_counts_as_public(self):
        self.assertEqual(self.visibility("<div>" + "x" * 70000 + "</div>"), "cong_khai")

    def test_short_page_or_network_error_is_reported_as_unknown(self):
        self.assertEqual(self.visibility("<div>trang ngắn</div>"), "khong_ro")
        self.assertEqual(self.visibility("", status=500), "khong_ro")

    def test_no_permalink_is_unknown(self):
        from plugins.zalo_tools import facebook as fb

        self.assertEqual(fb.public_visibility(""), "khong_ro")


class ZaloFbCheckTest(unittest.IsolatedAsyncioTestCase):
    async def test_check_reports_both_what_facebook_claims_and_what_outsiders_see(self):
        from plugins.zalo_tools import facebook as fb

        page = {"id": "301466423591522", "name": "Đoàn trường", "token": "x", "default": True}
        graph_result = {
            "id": "301466423591522_1363283005917157",
            "is_published": True, "is_hidden": False, "timeline_visibility": "normal",
            "privacy": {"description": "Công khai"},
            "permalink_url": "https://www.facebook.com/1361284372783687/posts/1363283005917157",
        }
        with patch.object(fb, "resolve_page", return_value=(page, "")), \
                patch.object(fb, "graph", return_value=graph_result), \
                patch.object(fb, "public_visibility", return_value="khong_xem_duoc"):
            out = json.loads(await zalo_tools.zalo_fb_check({"post_id": graph_result["id"]}))

        self.assertTrue(out["success"])
        result = out["result"]
        self.assertTrue(result["facebook_khai"]["da_dang"])
        self.assertEqual(result["facebook_khai"]["quyen"], "Công khai")
        self.assertEqual(result["hien_thi_cong_khai"], "khong_xem_duoc")
        self.assertIn("KHÔNG cho người ngoài xem", result["huong_dan"])

    async def test_check_requires_a_post_id(self):
        out = json.loads(await zalo_tools.zalo_fb_check({}))
        self.assertFalse(out["success"])
        self.assertIn("post_id", out["error"])


class ZaloKbScopeTest(unittest.IsolatedAsyncioTestCase):
    """ZALO_KB_PUBLIC_DIRS đóng phần còn lại của kho: kho thật là cả một ổ đĩa
    nhiều năm, người trong nhóm chỉ được thấy vài thư mục của năm hiện hành."""

    def setUp(self):
        self.root = tempfile.TemporaryDirectory()
        base = self.root.name
        for rel in ("ĐOÀN CNT 26-27/VĂN BẢN PHÁT RA 26-27/21-KH.txt",
                    "ĐOÀN CNT 23-24/23-24 VĂN BẢN PHÁT RA/01-KH.txt",
                    "ngay-goc.txt"):
            path = os.path.join(base, *rel.split("/"))
            os.makedirs(os.path.dirname(path), exist_ok=True)
            with open(path, "w", encoding="utf-8") as fh:
                fh.write("nội dung " + rel)
        zalo_tools._KB_CACHE.update(root=None, at=0.0, files=None, skipped=0)
        self.addCleanup(zalo_tools._KB_CACHE.update, root=None, at=0.0, files=None, skipped=0)
        self.addCleanup(self.root.cleanup)
        self.enterContext(patch("agent.secret_scope.get_secret",
                                side_effect=lambda name, default="": os.environ.get(name, default)))
        self.enterContext(patch.dict(os.environ, {"ZALO_KB_DIR": base}))

    def paths(self, payload):
        return sorted(f["path"] for f in json.loads(payload)["result"]["files"])

    async def test_public_dirs_hide_other_years_and_files_at_the_root(self):
        with patch.dict(os.environ, {"ZALO_KB_PUBLIC_DIRS": "ĐOÀN CNT 26-27, ĐOÀN CNT 25 - 26"}):
            self.assertEqual(
                self.paths(await zalo_tools.zalo_kb_list({})),
                ["ĐOÀN CNT 26-27/VĂN BẢN PHÁT RA 26-27/21-KH.txt"],
            )
            # Đoán đúng tên tệp ngoài phạm vi cũng không đọc được.
            denied = json.loads(await zalo_tools.zalo_kb_read(
                {"path": "ĐOÀN CNT 23-24/23-24 VĂN BẢN PHÁT RA/01-KH.txt"}))
            self.assertFalse(denied["success"])
            allowed = json.loads(await zalo_tools.zalo_kb_read(
                {"path": "ĐOÀN CNT 26-27/VĂN BẢN PHÁT RA 26-27/21-KH.txt"}))
            self.assertTrue(allowed["success"])
            self.assertIn("21-KH.txt", allowed["result"]["content"])

    async def test_large_store_lists_top_level_folders_so_the_agent_can_narrow_down(self):
        big = os.path.join(self.root.name, "ĐOÀN CNT 24-25")
        os.makedirs(big, exist_ok=True)
        for i in range(205):
            with open(os.path.join(big, f"tep-{i:03d}.txt"), "w", encoding="utf-8") as fh:
                fh.write("x")
        zalo_tools._KB_CACHE.update(root=None, at=0.0, files=None, skipped=0)

        with patch.dict(os.environ, {"ZALO_KB_PUBLIC_DIRS": ""}):
            payload = json.loads(await zalo_tools.zalo_kb_list({}))["result"]

        self.assertEqual(len(payload["files"]), 200)
        self.assertEqual(payload["count"], 208)
        self.assertEqual(payload["folders"]["ĐOÀN CNT 24-25"], 205)
        self.assertEqual(payload["folders"]["ĐOÀN CNT 26-27"], 1)
        self.assertEqual(payload["folders"]["ĐOÀN CNT 23-24"], 1)
        self.assertIn("folders", payload["note"])

    async def test_no_setting_keeps_the_whole_store_visible(self):
        zalo_tools._KB_CACHE.update(root=None, at=0.0, files=None, skipped=0)
        with patch.dict(os.environ, {"ZALO_KB_PUBLIC_DIRS": ""}):
            self.assertEqual(len(self.paths(await zalo_tools.zalo_kb_list({}))), 3)


class BridgeKeepaliveBoundsTest(unittest.TestCase):
    """Keepalive phải sống lâu hơn MỌI cửa sổ chờ, không chỉ cửa sổ ack.

    Lỗi 18/09/2026: ping timeout 180s < approval timeout 300s, nên một công cụ
    chờ người bấm approve/deny sẽ tự giết đường gửi của chính nó — prompt không
    tới Zalo, người dùng không thể trả lời, rồi retry gửi mỗi 2 giây vô hạn.
    Đây là một quan hệ số học giữa hai tệp, thứ không lộ ra trong test hành vi.
    """

    def test_keepalive_outlives_every_wait_window(self):
        for name in ("SLOW_ACK_TIMEOUT_SECONDS", "APPROVAL_WAIT_CEILING_SECONDS"):
            with self.subTest(window=name):
                self.assertGreater(
                    zalo_adapter.BRIDGE_PING_TIMEOUT_SECONDS,
                    getattr(zalo_adapter, name),
                    f"BRIDGE_PING_TIMEOUT_SECONDS phải lớn hơn {name}: "
                    "một kết nối đóng trước khi hết cửa sổ chờ thì ack không bao giờ về",
                )

    def test_approval_ceiling_still_matches_the_gateway(self):
        # Nếu upstream đổi _APPROVAL_TIMEOUT_SECONDS, test này phải đỏ chứ không
        # được để adapter âm thầm dùng số cũ.
        from gateway import run as gateway_run

        self.assertEqual(
            zalo_adapter.APPROVAL_WAIT_CEILING_SECONDS,
            gateway_run.GatewayRunner._APPROVAL_TIMEOUT_SECONDS,
        )


if __name__ == "__main__":
    unittest.main()
